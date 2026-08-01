import hashlib
import json
from pathlib import Path
import random
import time

import pandas as pd
import pytest
import openpyxl
from openpyxl import Workbook

from cashflow_indirect.pipeline import (
    finalize_run,
    get_run_status,
    prepare_run,
)


JOURNAL_HEADER = [
    "记账时间",
    "会计月",
    "凭证种类",
    "凭证编号",
    "编号",
    "业务说明",
    "一级编号",
    "一级科目",
    "科目编号",
    "科目名称",
    "借方发生额",
    "贷方发生额",
    "对方科目",
    "匹配类型",
]

BALANCE_HEADER = [
    "行类型",
    "科目编码",
    "科目名称",
    "科目级次",
    "上级科目编码",
    "科目性质",
    "所属科目编码",
    "辅助类型",
    "辅助编码",
    "辅助名称",
    "期初借方余额",
    "期初贷方余额",
    "本期借方发生额",
    "本期贷方发生额",
    "期末借方余额",
    "期末贷方余额",
]


def _line(
    code,
    name,
    nature,
    *,
    debit=0,
    credit=0,
    level1_code=None,
    level1_name=None,
    original_line_id=None,
):
    return {
        "code": str(code),
        "name": name,
        "nature": nature,
        "debit": debit,
        "credit": credit,
        "level1_code": str(level1_code or code),
        "level1_name": level1_name or name,
        "original_line_id": original_line_id,
    }


def _build_accounting_input(path: Path) -> None:
    journal = []
    accounts = {}
    voucher_no = 0

    def add_voucher(summary, lines, *, match_type="规则匹配"):
        nonlocal voucher_no
        voucher_no += 1
        for index, row in enumerate(lines, start=1):
            accounts[row["code"]] = row
            counterparties = "、".join(
                sorted(
                    {
                        other["level1_name"]
                        for other in lines
                        if other is not row
                    }
                )
            )
            journal.append(
                [
                    "2026-12-31",
                    "12月",
                    "记",
                    f"{voucher_no:03d}",
                    row["original_line_id"] or str(index),
                    summary,
                    row["level1_code"],
                    row["level1_name"],
                    row["code"],
                    row["name"],
                    row["debit"],
                    row["credit"],
                    counterparties,
                    match_type,
                ]
            )

    add_voucher(
        "计提存货跌价准备",
        [
            _line("670101", "资产减值损失", "损益", debit=10),
            _line("140599", "存货跌价准备", "资产", credit=10),
        ],
    )
    add_voucher(
        "红字冲回存货跌价准备",
        [
            _line("670101", "资产减值损失", "损益", debit=-1),
            _line("140599", "存货跌价准备", "资产", credit=-1),
        ],
    )
    add_voucher(
        "计提应收账款信用减值",
        [
            _line("670201", "信用减值损失", "损益", debit=8),
            _line("112299", "坏账准备", "资产", credit=8),
        ],
    )
    add_voucher(
        "核销存货及其跌价准备",
        [
            _line("140599", "存货跌价准备", "资产", debit=2),
            _line("140501", "库存商品", "资产", credit=2),
        ],
    )
    add_voucher(
        "核销无法收回应收账款",
        [
            _line("112299", "坏账准备", "资产", debit=3),
            _line("112201", "应收账款", "资产", credit=3),
        ],
    )
    add_voucher(
        "生产设备折旧计入存货成本",
        [
            _line("510101", "制造费用—折旧费", "成本", debit=7),
            _line("160299", "累计折旧", "资产", credit=7),
        ],
    )
    add_voucher(
        "无形资产摊销",
        [
            _line("660201", "管理费用—无形资产摊销", "损益", debit=6),
            _line("170299", "累计摊销", "资产", credit=6),
        ],
    )
    add_voucher(
        "长期待摊费用摊销",
        [
            _line("660202", "管理费用—长期待摊费用摊销", "损益", debit=5),
            _line("180101", "长期待摊费用", "资产", credit=5),
        ],
    )
    add_voucher(
        "出售固定资产净损失",
        [
            _line(
                "611501",
                "资产处置损益",
                "损益",
                debit=2,
                original_line_id="1",
            ),
            _line(
                "611501",
                "资产处置损益",
                "损益",
                debit=2,
                original_line_id="1",
            ),
            _line("160601", "固定资产清理", "资产", credit=4),
        ],
        match_type="算法生成(多借多贷拆分)",
    )
    add_voucher(
        "固定资产毁损报废净损失",
        [
            _line("671101", "营业外支出—固定资产报废损失", "损益", debit=3),
            _line("160601", "固定资产清理", "资产", credit=3),
        ],
    )
    add_voucher(
        "固定资产损失待判断",
        [
            _line("671102", "固定资产损失", "损益", debit=2),
            _line("160601", "固定资产清理", "资产", credit=2),
        ],
    )
    add_voucher(
        "交易性金融资产公允价值下降",
        [
            _line("610101", "公允价值变动损益", "损益", debit=2),
            _line("110101", "交易性金融资产", "资产", credit=2),
        ],
    )
    add_voucher(
        "借款利息费用",
        [
            _line("660301", "财务费用—借款利息", "损益", debit=3),
            _line("220201", "应付利息", "负债", credit=3),
        ],
    )
    add_voucher(
        "现金及现金等价物汇兑收益",
        [
            _line("100201", "银行存款", "资产", debit=1),
            _line("660302", "财务费用—现金汇兑损益", "损益", credit=1),
        ],
    )
    add_voucher(
        "确认权益法投资收益",
        [
            _line("151101", "长期股权投资", "资产", debit=2),
            _line("611101", "投资收益", "损益", credit=2),
        ],
    )
    add_voucher(
        "递延所得税资产减少计入损益",
        [
            _line("680101", "所得税费用", "损益", debit=1),
            _line("181101", "递延所得税资产", "资产", credit=1),
        ],
    )
    add_voucher(
        "递延所得税负债增加计入损益",
        [
            _line("680101", "所得税费用", "损益", debit=1),
            _line("290101", "递延所得税负债", "负债", credit=1),
        ],
    )
    add_voucher(
        "其他综合收益对应递延所得税负债增加",
        [
            _line("400101", "其他综合收益", "权益", debit=2),
            _line("290101", "递延所得税负债", "负债", credit=2),
        ],
    )
    add_voucher(
        "权益结算股份支付费用",
        [
            _line("660401", "管理费用—权益结算股份支付", "损益", debit=14),
            _line("400201", "资本公积", "权益", credit=14),
        ],
    )
    add_voucher(
        "政府补助A当期收款并全额计入损益",
        [
            _line("100201", "银行存款", "资产", debit=10),
            _line("611701", "其他收益—政府补助", "损益", credit=10),
        ],
    )
    add_voucher(
        "政府补助B收款大于当期损益",
        [
            _line("100201", "银行存款", "资产", debit=10),
            _line("611701", "其他收益—政府补助", "损益", credit=2),
            _line("240101", "递延收益—政府补助", "负债", credit=8),
        ],
    )
    add_voucher(
        "政府补助C本期摊销未收款",
        [
            _line("240101", "递延收益—政府补助", "负债", debit=2),
            _line("611701", "其他收益—政府补助", "损益", credit=2),
        ],
    )
    add_voucher(
        "政府补助D确认应收并计入损益",
        [
            _line("112301", "应收政府补助", "资产", debit=10),
            _line("611701", "其他收益—政府补助", "损益", credit=10),
        ],
    )
    add_voucher(
        "政府补助E收回应收款未形成损益",
        [
            _line("100201", "银行存款", "资产", debit=10),
            _line("112301", "应收政府补助", "资产", credit=10),
        ],
    )

    period_totals = {}
    for row in journal:
        code = str(row[8])
        debit, credit = row[10], row[11]
        current = period_totals.setdefault(code, [0, 0])
        current[0] += debit
        current[1] += credit

    balances = {
        "140501": (100, 0, 80, 0),
        "112201": (100, 0, 80, 0),
        "220101": (0, 30, 0, 45),
        "122101": (50, 0, 40, 0),
    }
    accounts["220101"] = _line("220101", "应付账款", "负债")
    accounts["122101"] = _line("122101", "其他应收款", "资产")
    balance_rows = []
    for code, row in accounts.items():
        opening_debit, opening_credit, closing_debit, closing_credit = balances.get(
            code, (0, 0, 0, 0)
        )
        period_debit, period_credit = period_totals.get(code, (0, 0))
        balance_rows.append(
            [
                "科目",
                code,
                row["name"],
                1,
                "",
                row["nature"],
                "",
                "",
                "",
                "",
                opening_debit,
                opening_credit,
                period_debit,
                period_credit,
                closing_debit,
                closing_credit,
            ]
        )
    balance_rows.extend(
        [
            [
                "辅助",
                "",
                "",
                "",
                "",
                "",
                "122101",
                "项目",
                "P01",
                "员工备用金",
                30,
                0,
                0,
                0,
                25,
                0,
            ],
            [
                "辅助",
                "",
                "",
                "",
                "",
                "",
                "122101",
                "项目",
                "P02",
                "投资款",
                20,
                0,
                0,
                0,
                15,
                0,
            ],
        ]
    )
    workbook = Workbook()
    journal_sheet = workbook.active
    journal_sheet.title = "序时账"
    for row in [JOURNAL_HEADER, *journal]:
        journal_sheet.append(row)
    balance_sheet = workbook.create_sheet("余额表")
    for row in [BALANCE_HEADER, *balance_rows]:
        balance_sheet.append(row)
    workbook.save(path)


def _write_config(tmp_path: Path, input_path: Path) -> Path:
    config = {
        "journal_path": str(input_path),
        "journal_sheet": "序时账",
        "trial_balance_path": str(input_path),
        "trial_balance_sheet": "余额表",
        "output_root": str(tmp_path / "runs"),
        "rules_path": str(Path("rules/cas_general_v1.json").resolve()),
        "entity_name": "端到端测试公司",
        "period_start": "2026-01-01",
        "period_end": "2026-12-31",
        "functional_currency": "CNY",
        "currency_scale": 2,
        "display_unit": "元",
        "net_profit_minor": 0,
        "journal_preprocessing_confirmed": True,
        "net_profit_user_confirmed": True,
        "net_profit_source": "user_provided",
        "net_profit_source_note": "端到端测试客户直接提供净利润",
        "performance_materiality_minor": 100,
        "direct_cfo_minor": 11600,
        "protocol_version": "1.0",
        "code_version": "e2e-test",
    }
    config_path = tmp_path / "运行配置.json"
    config_path.write_text(
        json.dumps(config, ensure_ascii=False, indent=2),
        encoding="utf-8-sig",
    )
    return config_path


def _write_ai_decision(request: dict, path: Path) -> None:
    other_item = next(item for item in request["candidate_items"] if item != 7)
    decision = {
        "request_id": request["request_id"],
        "recommended_item": 7,
        "recommended_status": "included",
        "recommended_amount_minor": request["available_amount_minor"],
        "recommendation": "作为固定资产报废损失计入第7项",
        "reason": "摘要只表明固定资产损失，结合营业外支出列报建议按报废处理",
        "supporting_evidence_ids": request["evidence_ids"],
        "contrary_evidence": ["未提供报废审批单，保留建议理由"],
        "rejected_alternatives": [
            {"item_no": other_item, "reason": "没有正常出售或转让的明确证据"}
        ],
        "applicable_rule_ids": [
            rule["rule_id"]
            for rule in request["candidate_rules"]
            if rule["item_no"] == 7
        ],
        "cfo_impact_minor": request["available_amount_minor"],
        "confidence": 0.82,
        "strong_conflict": False,
        "human_review_required": False,
        "model_id": "cashflow-indirect-ai",
        "prompt_template_version": "cashflow-ai-v1",
        "prompt_hash": "a" * 64,
        "response_hash": "b" * 64,
    }
    path.write_text(
        json.dumps(decision, ensure_ascii=False) + "\n",
        encoding="utf-8-sig",
    )


def _verify_jsonl(path: Path) -> bool:
    previous_hash = "0" * 64
    for expected_sequence, line in enumerate(
        path.read_text(encoding="utf-8-sig").splitlines(),
        start=1,
    ):
        event = json.loads(line)
        if event["sequence_no"] != expected_sequence:
            return False
        if event["previous_event_hash"] != previous_hash:
            return False
        event_hash = event.pop("event_hash")
        canonical = json.dumps(
            event,
            ensure_ascii=False,
            sort_keys=True,
            separators=(",", ":"),
        )
        if hashlib.sha256(canonical.encode("utf-8")).hexdigest() != event_hash:
            return False
        previous_hash = event_hash
    return True


def _verify_manifest_hashes(path: Path) -> bool:
    manifest = json.loads(path.read_text(encoding="utf-8-sig"))
    return all(
        hashlib.sha256((path.parent / item["path"]).read_bytes()).hexdigest()
        == item["sha256"]
        for item in manifest["files"]
    )


def _frozen_hashes(run_dir: Path) -> dict[str, str]:
    return {
        str(path.relative_to(run_dir)): hashlib.sha256(path.read_bytes()).hexdigest()
        for path in (
            run_dir / "02_normalized" / "journal_lines.parquet",
            run_dir / "02_normalized" / "trial_balance.parquet",
            run_dir / "04_candidates" / "candidate_scope.parquet",
            run_dir / "04_candidates" / "rule_claims.parquet",
        )
    }


def test_complete_chinese_workbook_reconciles_all_sixteen_items(tmp_path):
    input_path = tmp_path / "端到端输入.xlsx"
    _build_accounting_input(input_path)
    prepared = prepare_run(_write_config(tmp_path, input_path))
    requests = [
        json.loads(line)
        for line in Path(prepared["ai_requests_path"])
        .read_text(encoding="utf-8-sig")
        .splitlines()
    ]

    assert prepared["status"] == "awaiting_ai"
    assert len(requests) == 1

    ai_path = tmp_path / "ai结论.jsonl"
    _write_ai_decision(requests[0], ai_path)
    result = finalize_run(Path(prepared["run_dir"]), ai_path)
    summary = json.loads(
        (
            Path(prepared["run_dir"])
            / "08_output"
            / "result_summary.json"
        ).read_text(encoding="utf-8-sig")
    )
    reconciliation = json.loads(
        (
            Path(prepared["run_dir"])
            / "08_output"
            / "reconciliation.json"
        ).read_text(encoding="utf-8-sig")
    )

    assert {row["item_no"]: row["amount_minor"] for row in summary["items"]} == {
        1: 900,
        2: 800,
        3: 700,
        4: 600,
        5: 500,
        6: 400,
        7: 500,
        8: 200,
        9: 200,
        10: -200,
        11: 100,
        12: 100,
        13: 1100,
        14: 2200,
        15: 1500,
        16: 2000,
    }
    assert result["unclassified_amount_minor"] == 0
    assert result["exclusive_conflict_count"] == 0
    assert result["indirect_cfo_minor"] == 11600
    assert reconciliation["status"] == "pass"
    assert {
        check["check_id"]: check["difference_minor"]
        for check in reconciliation["checks"]
    } == {
        "voucher_balance": 0,
        "journal_trial_balance": 0,
        "amount_occupancy": 0,
        "exclusive_items": 0,
        "balance_bridge": 0,
        "deferred_tax": 0,
        "direct_indirect_cfo": 0,
    }
    run_dir = Path(prepared["run_dir"])
    assert _verify_jsonl(run_dir / "07_audit" / "events.jsonl")
    assert _verify_manifest_hashes(
        run_dir / "08_output" / "output_manifest.json"
    )
    manifest = json.loads(
        (run_dir / "08_output" / "output_manifest.json").read_text(
            encoding="utf-8-sig"
        )
    )
    manifest_paths = {entry["path"] for entry in manifest["files"]}
    assert {"events.jsonl", "event_index.parquet"} <= manifest_paths
    assert len(pd.read_parquet(run_dir / "08_output" / "event_index.parquet")) > 0
    workbook = openpyxl.load_workbook(
        result["workbook_path"],
        read_only=False,
    )
    try:
        # 余额桥接已并入各项目底稿顶部区域（政府补助桥接归入第16项）
        bridge_text = "\n".join(
            str(cell.value)
            for row in workbook["16其他明确调整"].iter_rows()
            for cell in row
            if cell.value is not None
        )
        assert bridge_text.count("政府补助") >= 1
    finally:
        workbook.close()
    normalized = pd.read_parquet(
        run_dir / "02_normalized" / "journal_lines.parquet"
    )
    split = normalized.loc[
        normalized["summary"].eq("出售固定资产净损失")
        & normalized["original_line_id"].eq("1")
    ]
    assert split["split_index"].tolist() == [0, 1]


def test_restart_reuses_checkpoint_and_source_change_invalidates_it(tmp_path):
    input_path = tmp_path / "检查点输入.xlsx"
    _build_accounting_input(input_path)
    prepared = prepare_run(_write_config(tmp_path, input_path))
    run_dir = Path(prepared["run_dir"])
    before = _frozen_hashes(run_dir)

    assert get_run_status(run_dir)["status"] == "awaiting_ai"
    assert _frozen_hashes(run_dir) == before

    request = json.loads(
        Path(prepared["ai_requests_path"])
        .read_text(encoding="utf-8-sig")
        .splitlines()[0]
    )
    ai_path = tmp_path / "检查点AI结论.jsonl"
    _write_ai_decision(request, ai_path)
    assert finalize_run(run_dir, ai_path)["status"] == "exported"
    assert _frozen_hashes(run_dir) == before

    input_path.write_bytes(input_path.read_bytes() + b"source-changed")
    status = get_run_status(run_dir)
    assert status["status"] == "invalidated"
    assert status["invalidated_inputs"] == ["journal", "trial_balance"]
    with pytest.raises(ValueError, match="旧检查点失效"):
        finalize_run(run_dir, ai_path)


@pytest.mark.performance
def test_prepare_one_hundred_thousand_journal_rows_within_two_minutes(tmp_path):
    generator = random.Random(20260731)
    workbook = Workbook(write_only=True)
    journal_sheet = workbook.create_sheet("序时账")
    journal_sheet.append(JOURNAL_HEADER)
    total = 0
    for voucher in range(1, 50_001):
        amount = generator.randint(1, 99)
        total += amount
        journal_sheet.append(
            [
                "2026-12-31",
                "12月",
                "记",
                f"{voucher:05d}",
                "1",
                "支付经营性往来款",
                "1002",
                "银行存款",
                "100201",
                "银行存款",
                amount,
                0,
                "应付账款",
                "规则匹配",
            ]
        )
        journal_sheet.append(
            [
                "2026-12-31",
                "12月",
                "记",
                f"{voucher:05d}",
                "2",
                "支付经营性往来款",
                "2201",
                "应付账款",
                "220101",
                "应付账款",
                0,
                amount,
                "银行存款",
                "规则匹配",
            ]
        )
    balance_sheet = workbook.create_sheet("余额表")
    balance_sheet.append(BALANCE_HEADER)
    balance_sheet.append(
        [
            "科目",
            "100201",
            "银行存款",
            1,
            "",
            "资产",
            "",
            "",
            "",
            "",
            0,
            0,
            total,
            0,
            total,
            0,
        ]
    )
    balance_sheet.append(
        [
            "科目",
            "220101",
            "应付账款",
            1,
            "",
            "负债",
            "",
            "",
            "",
            "",
            0,
            0,
            0,
            total,
            0,
            total,
        ]
    )
    input_path = tmp_path / "十万行输入.xlsx"
    workbook.save(input_path)
    config_path = _write_config(tmp_path, input_path)
    config = json.loads(config_path.read_text(encoding="utf-8-sig"))
    config["direct_cfo_minor"] = 0
    config_path.write_text(
        json.dumps(config, ensure_ascii=False, indent=2),
        encoding="utf-8-sig",
    )

    started = time.perf_counter()
    result = prepare_run(config_path)
    elapsed = time.perf_counter() - started
    state = json.loads(
        (Path(result["run_dir"]) / "run_state.json").read_text(
            encoding="utf-8-sig"
        )
    )

    assert result["status"] == "ready_to_finalize"
    assert result["ai_request_count"] == 0
    assert elapsed <= 120
    assert state["timings_seconds"]["prepare"] <= 120
    assert state["peak_memory_bytes"] > 0


def test_decision_trace_and_item_sheets_keep_reason_and_evidence(tmp_path):
    import openpyxl

    input_path = tmp_path / "证据端到端输入.xlsx"
    _build_accounting_input(input_path)
    prepared = prepare_run(_write_config(tmp_path, input_path))
    requests = [
        json.loads(line)
        for line in Path(prepared["ai_requests_path"])
        .read_text(encoding="utf-8-sig")
        .splitlines()
    ]
    assert len(requests) == 1
    ai_path = tmp_path / "ai结论.jsonl"
    _write_ai_decision(requests[0], ai_path)
    result = finalize_run(Path(prepared["run_dir"]), ai_path)

    assert result["status"] == "exported"
    run_dir = Path(prepared["run_dir"])
    trace = pd.read_parquet(run_dir / "06_decisions" / "decision_trace.parquet")
    assert not trace.empty
    # 规则轨迹保留理由与占用
    rule_rows = trace.loc[trace["item_no"].eq(1)]
    assert not rule_rows.empty
    assert rule_rows["decision_reason"].fillna("").astype(str).str.len().gt(0).all()
    assert (rule_rows["physical_consumed_minor"] > 0).all()
    # AI 候选改归轨迹保留建议理由、正反证据和未采用方案
    ai_rows = trace.loc[trace["candidate_status"].eq("ai_candidate")]
    assert not ai_rows.empty
    assert ai_rows["recommendation_reason"].fillna("").astype(str).str.len().gt(0).all()
    assert "报废审批单" in "".join(
        ai_rows["contrary_evidence"].fillna("").astype(str).tolist()
    )
    assert ai_rows["rejected_alternatives"].fillna("").astype(str).str.len().gt(0).all()
    # 项目底稿同时展示判断证据
    workbook = openpyxl.load_workbook(
        run_dir / "08_output" / "现金流量表附表计算结果.xlsx",
        read_only=False,
    )
    try:
        item1 = workbook["01资产减值准备"]
        item1_text = "\n".join(
            str(cell.value)
            for row in item1.iter_rows()
            for cell in row
            if cell.value is not None
        )
        assert "计入当期损益" in item1_text or "减值" in item1_text
        item7 = workbook["07固定资产报废损失"]
        item7_text = "\n".join(
            str(cell.value)
            for row in item7.iter_rows()
            for cell in row
            if cell.value is not None
        )
        assert "报废" in item7_text
        assert "未采用" in item7_text or "备选" in item7_text
    finally:
        workbook.close()
