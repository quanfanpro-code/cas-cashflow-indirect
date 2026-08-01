import json
from pathlib import Path
from types import SimpleNamespace

import openpyxl
import pandas as pd
import pytest

import cashflow_indirect.output as output_module
from cashflow_indirect.output import OutputExportError, export_results
from cashflow_indirect.reconciliation import calculate_summary, reconcile_all


@pytest.fixture
def completed_result(tmp_path, rule_book):
    occupancy = pd.DataFrame(
        [
            {
                "fragment_id": "f-1",
                "item_no": 1,
                "adjustment_amount_minor": 20000,
                "claimed_amount_minor": 20000,
                "physical_consumed_minor": 20000,
                "remaining_amount_minor": 0,
                "decision": "included",
                "rule_id": "IMP-ASSET",
                "paired_offset_id": None,
            },
            {
                "fragment_id": "f-9",
                "item_no": 9,
                "adjustment_amount_minor": 30000,
                "claimed_amount_minor": 30000,
                "physical_consumed_minor": 30000,
                "remaining_amount_minor": 0,
                "decision": "ai_classified",
                "rule_id": "FIN-CASH-FX",
                "paired_offset_id": None,
            },
        ]
    )
    summary = calculate_summary(1_000_000, occupancy, rule_book)
    reconciliation = reconcile_all(summary, direct_cfo_minor=1_050_000)
    details = pd.DataFrame(
        [
            {
                "item_no": 1,
                "item_name": "资产减值准备",
                "business_event_id": "event-1",
                "fragment_id": "f-1",
                "voucher_date": "2026-12-31",
                "voucher_no": "记-001",
                "source_row": 2,
                "account_name": "资产减值损失",
                "counterparty_level1": "存货跌价准备",
                "summary": "计提跌价准备",
                "source_amount_minor": 20000,
                "adjustment_amount_minor": 20000,
                "decision": "included",
                "rule_id": "IMP-ASSET",
                "reason": "计入当期损益的减值",
            }
        ]
    )
    bridges = pd.DataFrame(
        [
            {
                "item_no": 13,
                "record_type": "summary",
                "opening_minor": 500000,
                "closing_minor": 450000,
                "base_change_minor": 50000,
                "corrections_minor": 0,
                "final_amount_minor": 50000,
                "reason": "存货原值余额桥接",
            }
        ]
    )
    issues = pd.DataFrame(
        [
            {
                "issue_code": "ACCOUNTING_WARNING",
                "amount_minor": 1000,
                "item_no": 3,
                "reason": "示例会计异常",
                "recommendation": "复核折旧政策",
                "recommendation_reason": "资产状态与折旧政策不一致",
                "supporting_evidence": "固定资产卡片",
                "contrary_evidence": "无",
                "alternatives": "按实际损益调整",
                "review_required": False,
                "final_decision": "保留调整并提示异常",
            }
        ]
    )
    output_dir = tmp_path / "output"
    run_root = tmp_path / "run"
    run_root.mkdir()
    run_store = SimpleNamespace(
        root=run_root,
        run_id="run-test",
        input_fingerprint="input-hash",
        rule_version=rule_book.rule_version,
        code_version="test-code",
    )
    journal_snapshot = tmp_path / "序时账快照.xlsx"
    journal_workbook = openpyxl.Workbook()
    journal_worksheet = journal_workbook.active
    journal_worksheet.title = "序时账"
    journal_worksheet.append(
        [
            "记账时间", "会计月", "凭证种类", "凭证编号", "编号",
            "业务说明", "科目编号", "科目名称", "借方发生额", "贷方发生额",
            "对方科目", "匹配类型",
        ]
    )
    journal_worksheet.append(
        [
            "2026-12-31", "12月", "记", "001", "1",
            "计提资产减值准备", "671101", "资产减值损失",
            20000, 0, "存货跌价准备", "算法生成",
        ]
    )
    journal_worksheet.append(
        [
            "2026-12-31", "12月", "记", "001", "2",
            "结转损益", "410401", "本年利润",
            0, 20000, "资产减值损失", "结转损益规则",
        ]
    )
    journal_workbook.save(journal_snapshot)
    journal_workbook.close()
    decision_trace = pd.DataFrame(
        [
            {
                "trace_id": "trace-1",
                "fragment_id": "f-1",
                "business_event_id": "event-1",
                "source_row": 2,
                "item_no": 1,
                "candidate_status": "rule_classified",
                "final_decision": "included",
                "rule_id": "IMP-ASSET",
                "decision_reason": "计入当期损益的减值",
                "confidence": 1.0,
                "recommendation": "",
                "recommendation_reason": "",
                "supporting_evidence": "",
                "contrary_evidence": "未发现相反证据",
                "rejected_alternatives": "",
                "human_review_required": False,
                "claim_amount_minor": 20000,
                "physical_consumed_minor": 20000,
                "remaining_amount_minor": 0,
                "occupancy_type": "exclusive",
                "paired_offset_id": None,
            }
        ],
        columns=[
            "trace_id", "fragment_id", "business_event_id", "source_row",
            "item_no", "candidate_status", "final_decision", "rule_id",
            "decision_reason", "confidence", "recommendation",
            "recommendation_reason", "supporting_evidence",
            "contrary_evidence", "rejected_alternatives",
            "human_review_required", "claim_amount_minor",
            "physical_consumed_minor", "remaining_amount_minor",
            "occupancy_type", "paired_offset_id",
        ],
    )
    return {
        "run_store": run_store,
        "summary": summary,
        "details": details,
        "occupancy": occupancy,
        "bridges": bridges,
        "issues": issues,
        "reconciliation": reconciliation,
        "decision_trace": decision_trace,
        "journal_snapshot": journal_snapshot,
        "journal_sheet": "序时账",
        "journal_header_row": 1,
        "run_info": {
            "entity_name": "甲公司",
            "period_start": "2026-01-01",
            "period_end": "2026-12-31",
            "functional_currency": "CNY",
            "display_unit": "元",
            "display_divisor_minor": 100,
            "performance_materiality_minor": 100000,
            "protocol_version": "1.0",
            "run_status": "completed",
        },
        "output_dir": output_dir,
    }


def test_output_formats_share_identical_amounts(completed_result):
    output_dir = completed_result["output_dir"]

    manifest = export_results(**completed_result)

    json_summary = json.loads(
        (output_dir / "result_summary.json").read_text("utf-8-sig")
    )
    parquet_summary = pd.read_parquet(output_dir / "adjustment_summary.parquet")
    workbook = openpyxl.load_workbook(
        output_dir / "现金流量表附表计算结果.xlsx",
        data_only=True,
    )
    excel_values = {
        row[0].value: row[1].value
        for row in workbook["附表结果"].iter_rows(min_row=2)
        if row[0].value
    }

    assert json_summary["indirect_cfo_minor"] == (
        json_summary["net_profit_minor"]
        + parquet_summary["amount_minor"].sum()
    )
    assert (
        excel_values["经营活动产生的现金流量净额"]
        == json_summary["indirect_cfo_minor"] / 100
    )
    assert all((output_dir / entry["path"]).exists() for entry in manifest["files"])


def test_workbook_has_fixed_eighteen_sheets_and_manifest_is_portable(
    completed_result,
    rule_book,
):
    output_dir = completed_result["output_dir"]

    manifest = export_results(**completed_result)

    workbook = openpyxl.load_workbook(
        output_dir / "现金流量表附表计算结果.xlsx",
        read_only=True,
    )
    expected_sheets = ["附表结果", "全量序时账"] + [
        f"{item.item_no:02d}{item.name}" for item in rule_book.items
    ]
    assert workbook.sheetnames == expected_sheets
    assert all(not Path(entry["path"]).is_absolute() for entry in manifest["files"])
    assert all(len(entry["sha256"]) == 64 for entry in manifest["files"])
    assert {
        entry["path"] for entry in manifest["files"]
    } >= {
        "现金流量表附表计算结果.xlsx",
        "result_summary.json",
        "adjustment_summary.parquet",
        "adjustment_details.parquet",
        "amount_occupancy.parquet",
        "reconciliation.json",
        "issues.jsonl",
        "decision_trace.parquet",
        "output_status.json",
    }
    workbook_entry = next(
        entry
        for entry in manifest["files"]
        if entry["path"] == "现金流量表附表计算结果.xlsx"
    )
    assert workbook_entry["record_count"] == 18
    trace_entry = next(
        entry
        for entry in manifest["files"]
        if entry["path"] == "decision_trace.parquet"
    )
    assert len(trace_entry["sha256"]) == 64


def test_excel_failure_preserves_machine_outputs_and_can_retry(
    completed_result,
    monkeypatch,
):
    output_dir = completed_result["output_dir"]
    original = output_module.write_audit_workbook

    def occupied_workbook(*args, **kwargs):
        raise PermissionError("目标工作簿被占用")

    monkeypatch.setattr(output_module, "write_audit_workbook", occupied_workbook)
    with pytest.raises(OutputExportError, match="Excel"):
        export_results(**completed_result)

    assert (output_dir / "adjustment_summary.parquet").exists()
    failed_status = json.loads(
        (output_dir / "output_status.json").read_text("utf-8-sig")
    )
    assert failed_status["status"] == "failed"
    machine_hash = output_module.sha256_file(
        output_dir / "adjustment_summary.parquet"
    )

    monkeypatch.setattr(output_module, "write_audit_workbook", original)
    manifest = export_results(**completed_result)

    assert output_module.sha256_file(
        output_dir / "adjustment_summary.parquet"
    ) == machine_hash
    assert json.loads(
        (output_dir / "output_status.json").read_text("utf-8-sig")
    )["status"] == "completed"
    assert manifest["status"] == "completed"



def test_excel_value_error_marks_failed_and_preserves_machine_outputs(
    completed_result,
    monkeypatch,
):
    output_dir = completed_result["output_dir"]
    original = output_module.write_audit_workbook

    def overflow_workbook(*args, **kwargs):
        raise ValueError("输入序时账行数超出单表行数上限")

    monkeypatch.setattr(
        output_module, "write_audit_workbook", overflow_workbook
    )
    with pytest.raises(OutputExportError, match="机器结果已保留"):
        export_results(**completed_result)

    failed_status = json.loads(
        (output_dir / "output_status.json").read_text("utf-8-sig")
    )
    assert failed_status["status"] == "failed"
    assert failed_status["retryable"] is True
    assert "超出单表行数上限" in failed_status["error"]
    assert (output_dir / "adjustment_summary.parquet").exists()

    monkeypatch.setattr(output_module, "write_audit_workbook", original)



def test_rapid_mode_exports_two_sheets_and_sixteen_markdowns(
    completed_result,
    rule_book,
):
    output_dir = completed_result["output_dir"]

    manifest = export_results(**completed_result, output_mode="rapid")

    workbook = openpyxl.load_workbook(
        output_dir / "现金流量表附表计算结果.xlsx",
        read_only=True,
    )
    try:
        assert workbook.sheetnames == ["附表结果", "简报"]
    finally:
        workbook.close()
    workbook_entry = next(
        entry
        for entry in manifest["files"]
        if entry["path"] == "现金流量表附表计算结果.xlsx"
    )
    assert workbook_entry["record_count"] == 2
    md_entries = [
        entry
        for entry in manifest["files"]
        if entry["path"].startswith("底稿MD/") and entry["path"].endswith(".md")
    ]
    assert len(md_entries) == 16
    for item in rule_book.items:
        assert any(
            entry["path"] == f"底稿MD/{item.item_no:02d}{item.name}.md"
            for entry in md_entries
        )
