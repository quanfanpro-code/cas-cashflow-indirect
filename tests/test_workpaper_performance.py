"""一万行流式输出性能基线与峰值内存记录。"""

import os
from pathlib import Path
import time
import tracemalloc
from types import SimpleNamespace

import openpyxl
import pandas as pd
import pytest

from cashflow_indirect.reconciliation import calculate_summary, reconcile_all
from cashflow_indirect.workpaper import write_audit_workbook

PERF_ROW_COUNT = 10_000


@pytest.mark.performance
def test_ten_thousand_row_streaming_output_records_baseline(tmp_path, rule_book):
    """生成器构造一万行含扩展列的序时账，只测工作簿输出层。"""
    snapshot = tmp_path / "一万行序时账.xlsx"
    workbook = openpyxl.Workbook(write_only=True)
    worksheet = workbook.create_sheet("序时账")
    header = [
        "记账时间", "会计月", "凭证种类", "凭证编号", "编号", "业务说明",
        "科目编号", "科目名称", "借方发生额", "贷方发生额", "对方科目",
        "匹配类型", "核算维度", "部门", "员工", "银行账户", "物料",
    ]
    worksheet.append(header)
    for index in range(PERF_ROW_COUNT):
        worksheet.append(
            [
                "2026-12-31", "12月", "记", f"{index:04d}", str(index + 1),
                "测试业务说明行", "660201", "办公费",
                100, 0, "银行存款", "算法生成",
                "费用", "行政部", "员工甲", "账户A", "物料B",
            ]
        )
    workbook.save(snapshot)
    workbook.close()

    occupancy = pd.DataFrame(
        [
            {
                "fragment_id": "f-1",
                "item_no": 1,
                "adjustment_amount_minor": 1000000,
                "claimed_amount_minor": 1000000,
                "physical_consumed_minor": 1000000,
                "remaining_amount_minor": 0,
                "decision": "included",
                "rule_id": "IMP-ASSET",
                "paired_offset_id": None,
            }
        ]
    )
    summary = calculate_summary(1_000_000, occupancy, rule_book)
    reconciliation = reconcile_all(summary, direct_cfo_minor=2_000_000)
    details = pd.DataFrame(
        [
            {
                "item_no": 1,
                "item_name": "资产减值准备",
                "business_event_id": "event-1",
                "fragment_id": "f-1",
                "voucher_date": "2026-12-31",
                "voucher_no": "记-001",
                "source_row": 2,
                "account_name": "资产减值损失",
                "counterparty_level1": "存货跌价准备",
                "summary": "计提跌价准备",
                "source_amount_minor": 1000000,
                "adjustment_amount_minor": 1000000,
                "decision": "included",
                "rule_id": "IMP-ASSET",
                "reason": "计入当期损益的减值",
            }
        ]
    )
    bridges = pd.DataFrame()
    issues = pd.DataFrame()
    decision_trace = pd.DataFrame(
        [
            {
                "trace_id": "trace-1",
                "fragment_id": "f-1",
                "business_event_id": "event-1",
                "source_row": 2,
                "item_no": 1,
                "candidate_status": "rule_classified",
                "final_decision": "included",
                "rule_id": "IMP-ASSET",
                "decision_reason": "计入当期损益的减值",
                "confidence": 1.0,
                "recommendation": "",
                "recommendation_reason": "",
                "supporting_evidence": "",
                "contrary_evidence": "未发现相反证据",
                "rejected_alternatives": "",
                "human_review_required": False,
                "claim_amount_minor": 1000000,
                "physical_consumed_minor": 1000000,
                "remaining_amount_minor": 0,
                "occupancy_type": "exclusive",
                "paired_offset_id": None,
            }
        ],
        columns=[
            "trace_id", "fragment_id", "business_event_id", "source_row",
            "item_no", "candidate_status", "final_decision", "rule_id",
            "decision_reason", "confidence", "recommendation",
            "recommendation_reason", "supporting_evidence",
            "contrary_evidence", "rejected_alternatives",
            "human_review_required", "claim_amount_minor",
            "physical_consumed_minor", "remaining_amount_minor",
            "occupancy_type", "paired_offset_id",
        ],
    )
    run_info = {
        "entity_name": "性能测试主体",
        "period_start": "2026-01-01",
        "period_end": "2026-12-31",
        "functional_currency": "CNY",
        "display_unit": "元",
        "display_divisor_minor": 100,
        "performance_materiality_minor": 50000,
        "protocol_version": "1.0",
        "run_status": "exported",
    }
    target = tmp_path / "性能底稿.xlsx"

    tracemalloc.start()
    started = time.perf_counter()
    result = write_audit_workbook(
        target=target,
        summary=summary,
        details=details,
        bridges=bridges,
        issues=issues,
        reconciliation=reconciliation,
        run_info=run_info,
        decision_trace=decision_trace,
        journal_snapshot=snapshot,
        journal_sheet="序时账",
        journal_header_row=1,
    )
    elapsed = time.perf_counter() - started
    _, peak = tracemalloc.get_traced_memory()
    tracemalloc.stop()

    assert result["journal_row_count"] == PERF_ROW_COUNT
    assert result["link_count"] == 32
    size_bytes = Path(result["path"]).stat().st_size
    record = {
        "row_count": PERF_ROW_COUNT,
        "elapsed_seconds": round(elapsed, 3),
        "peak_memory_mb": round(peak / 1024 / 1024, 2),
        "size_bytes": size_bytes,
        "output_path": str(target),
    }
    print("PERF_BASELINE=" + str(record))
    assert elapsed < 600, f"一万行输出耗时过长：{elapsed:.1f}秒"
    # 输出文件可重新打开：18 张表，全量序时账表完整包含一万行
    workbook = openpyxl.load_workbook(target, read_only=True)
    try:
        assert len(workbook.sheetnames) == 18
        assert "全量序时账" in workbook.sheetnames
        full_sheet = workbook["全量序时账"]
        data_count = 0
        for _ in full_sheet.iter_rows(min_row=3):
            data_count += 1
        assert data_count == PERF_ROW_COUNT
    finally:
        workbook.close()
