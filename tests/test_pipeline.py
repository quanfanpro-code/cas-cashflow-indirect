import json
from pathlib import Path
import shutil

import openpyxl
import pandas as pd
import pytest

from cashflow_indirect.input_adapter import InputProtocolError
from cashflow_indirect.pipeline import (
    _account_group_mask,
    _asset_class,
    finalize_run,
    get_run_status,
    prepare_run,
)


def test_account_group_matches_nested_balance_path():
    frame = pd.DataFrame(
        {"account_path": ["资产/流动资产/库存商品", "资产/固定资产"]}
    )

    assert _account_group_mask(frame, ("库存商品",)).tolist() == [True, False]


@pytest.mark.parametrize(
    ("text", "expected"),
    [
        ("商誉减值准备", "goodwill"),
        ("长期股权投资减值准备", "long_term_equity_investment"),
        ("投资性房地产减值准备", "investment_property_cost"),
        ("在建工程减值准备", "construction_in_progress"),
        ("工程物资减值准备", "construction_materials"),
        ("持有待售资产减值准备", "held_for_sale_asset"),
    ],
)
def test_asset_class_covers_rule_book_impairment_assets(text, expected):
    assert _asset_class(text) == expected


def test_intangible_asset_scrap_is_item_six_not_fixed_asset_item_seven(
    minimal_config,
):
    config = json.loads(Path(minimal_config).read_text(encoding="utf-8-sig"))
    workbook = openpyxl.load_workbook(config["journal_path"])
    journal = workbook[config["journal_sheet"]]
    journal.cell(row=2, column=6, value="无形资产报废处置损失")
    journal.cell(row=2, column=8, value="资产处置损益")
    journal.cell(row=2, column=10, value="资产处置损益—无形资产报废")
    journal.cell(row=2, column=13, value="无形资产")
    journal.cell(row=3, column=6, value="无形资产报废处置损失")
    journal.cell(row=3, column=7, value="1701")
    journal.cell(row=3, column=8, value="无形资产")
    journal.cell(row=3, column=9, value="170101")
    journal.cell(row=3, column=10, value="无形资产")
    journal.cell(row=3, column=13, value="资产处置损益")
    balance = workbook[config["trial_balance_sheet"]]
    balance.cell(row=2, column=2, value="611501")
    balance.cell(row=2, column=3, value="资产处置损益—无形资产报废")
    balance.cell(row=3, column=2, value="170101")
    balance.cell(row=3, column=3, value="无形资产")
    workbook.save(config["journal_path"])
    workbook.close()

    prepared = prepare_run(minimal_config)
    claims = pd.read_parquet(
        Path(prepared["run_dir"]) / "04_candidates" / "rule_claims.parquet"
    )

    assert 6 in set(claims["item_no"])
    assert 7 not in set(claims["item_no"])


def write_minimal_input(path: Path) -> None:
    journal_rows = [
        [
            "记账时间",
            "会计月",
            "凭证种类",
            "凭证编号",
            "编号",
            "业务说明",
            "一级编号",
            "一级科目",
            "科目编号",
            "科目名称",
            "借方发生额",
            "贷方发生额",
            "对方科目",
            "匹配类型",
        ],
        [
            "2026-12-31",
            "12月",
            "记",
            "001",
            "1",
            "处置或报废固定资产待判断",
            "6711",
            "营业外支出",
            "671101",
            "固定资产损失",
            100,
            0,
            "固定资产清理",
            "算法生成(兜底拆分)",
        ],
        [
            "2026-12-31",
            "12月",
            "记",
            "001",
            "2",
            "处置或报废固定资产待判断",
            "1606",
            "固定资产清理",
            "160601",
            "固定资产清理",
            0,
            100,
            "营业外支出",
            "算法生成(兜底拆分)",
        ],
    ]
    balance_rows = [
        [
            "行类型",
            "科目编码",
            "科目名称",
            "科目级次",
            "上级科目编码",
            "科目性质",
            "期初借方余额",
            "期初贷方余额",
            "本期借方发生额",
            "本期贷方发生额",
            "期末借方余额",
            "期末贷方余额",
        ],
        ["科目", "671101", "固定资产损失", 1, "", "损益", 0, 0, 100, 0, 0, 0],
        ["科目", "160601", "固定资产清理", 1, "", "资产", 0, 0, 0, 100, 0, 100],
    ]
    workbook = openpyxl.Workbook()
    journal_sheet = workbook.active
    journal_sheet.title = "序时账"
    for row in journal_rows:
        journal_sheet.append(row)
    balance_sheet = workbook.create_sheet("余额表")
    for row in balance_rows:
        balance_sheet.append(row)
    workbook.save(path)


@pytest.fixture
def minimal_config(tmp_path):
    input_path = tmp_path / "最小输入.xlsx"
    write_minimal_input(input_path)
    config_path = tmp_path / "运行配置.json"
    config = {
        "journal_path": str(input_path),
        "journal_sheet": "序时账",
        "trial_balance_path": str(input_path),
        "trial_balance_sheet": "余额表",
        "output_root": str(tmp_path / "runs"),
        "rules_path": str(Path("rules/cas_general_v1.json").resolve()),
        "entity_name": "甲公司",
        "period_start": "2026-01-01",
        "period_end": "2026-12-31",
        "functional_currency": "CNY",
        "currency_scale": 2,
        "display_unit": "元",
        "net_profit_minor": -10000,
        "journal_preprocessing_confirmed": True,
        "net_profit_user_confirmed": True,
        "net_profit_source": "user_provided",
        "net_profit_source_note": "测试客户直接提供净利润",
        "performance_materiality_minor": 50000,
        "financial_entity_flag": False,
        "direct_cfo_minor": 0,
        "protocol_version": "1.0",
        "code_version": "test-code",
    }
    config_path.write_text(
        json.dumps(config, ensure_ascii=False, indent=2),
        encoding="utf-8-sig",
    )
    return config_path


def _update_config(config_path: Path, **updates) -> dict:
    config = json.loads(config_path.read_text(encoding="utf-8-sig"))
    config.update(updates)
    config_path.write_text(
        json.dumps(config, ensure_ascii=False, indent=2),
        encoding="utf-8-sig",
    )
    return config


def test_prepare_rejects_unconfirmed_journal_before_creating_run_directory(
    minimal_config,
):
    config = json.loads(minimal_config.read_text(encoding="utf-8-sig"))
    config.pop("journal_preprocessing_confirmed")
    minimal_config.write_text(
        json.dumps(config, ensure_ascii=False, indent=2),
        encoding="utf-8-sig",
    )

    with pytest.raises(InputProtocolError, match="序时账分析器"):
        prepare_run(minimal_config)

    assert not Path(config["output_root"]).exists()


def test_prepare_rejects_net_profit_not_confirmed_by_customer(minimal_config):
    _update_config(
        minimal_config,
        journal_preprocessing_confirmed=True,
        net_profit_user_confirmed=False,
    )

    with pytest.raises(InputProtocolError, match="客户.*净利润"):
        prepare_run(minimal_config)


def test_prepare_rejects_journal_calculated_net_profit(minimal_config):
    _update_config(
        minimal_config,
        journal_preprocessing_confirmed=True,
        net_profit_user_confirmed=True,
        net_profit_source="journal_calculated",
        net_profit_source_note="根据序时账损益科目推算",
    )

    with pytest.raises(InputProtocolError, match="序时账推算"):
        prepare_run(minimal_config)


def test_prepare_rejects_empty_net_profit_source_note(minimal_config):
    _update_config(
        minimal_config,
        journal_preprocessing_confirmed=True,
        net_profit_user_confirmed=True,
        net_profit_source="user_provided",
        net_profit_source_note="  ",
    )

    with pytest.raises(InputProtocolError, match="来源说明"):
        prepare_run(minimal_config)


def test_prepare_records_customer_confirmed_net_profit_source(minimal_config):
    _update_config(
        minimal_config,
        journal_preprocessing_confirmed=True,
        net_profit_user_confirmed=True,
        net_profit_source="user_confirmed_statement",
        net_profit_source_note="客户确认采用利润表本期净利润",
    )

    result = prepare_run(minimal_config)
    manifest = json.loads(
        (Path(result["run_dir"]) / "run_manifest.json").read_text(
            encoding="utf-8-sig"
        )
    )
    config_snapshot = json.loads(
        (
            Path(result["run_dir"])
            / "01_input_snapshot"
            / "config_snapshot.json"
        ).read_text(encoding="utf-8-sig")
    )

    expected_confirmation = {
        "journal_preprocessing_confirmed": True,
        "net_profit_user_confirmed": True,
        "net_profit_source": "user_confirmed_statement",
        "net_profit_source_note": "客户确认采用利润表本期净利润",
    }
    assert {
        field: manifest["spec"][field] for field in expected_confirmation
    } == expected_confirmation
    assert config_snapshot["net_profit_source_note"] == (
        "客户确认采用利润表本期净利润"
    )


def test_legacy_run_without_source_confirmation_is_invalidated(minimal_config):
    result = prepare_run(minimal_config)
    state_path = Path(result["run_dir"]) / "run_state.json"
    state = json.loads(state_path.read_text(encoding="utf-8-sig"))
    for field in (
        "journal_preprocessing_confirmed",
        "net_profit_user_confirmed",
        "net_profit_source",
        "net_profit_source_note",
    ):
        state["spec"].pop(field)
    state_path.write_text(
        json.dumps(state, ensure_ascii=False, indent=2),
        encoding="utf-8-sig",
    )

    status = get_run_status(Path(result["run_dir"]))

    assert status["status"] == "invalidated"
    assert status["invalidated_inputs"] == ["input_source_confirmation"]
    assert "序时账分析器" in status["invalidation_reason"]
    with pytest.raises(InputProtocolError, match="序时账分析器"):
        finalize_run(Path(result["run_dir"]), None)


def valid_ai_file(run_result, tmp_path):
    request_path = Path(run_result["ai_requests_path"])
    request = json.loads(request_path.read_text("utf-8-sig").splitlines()[0])
    other_item = next(item for item in request["candidate_items"] if item != 7)
    decision = {
        "request_id": request["request_id"],
        "recommended_item": 7,
        "recommended_status": "included",
        "recommended_amount_minor": request["available_amount_minor"],
        "recommendation": "作为固定资产报废损失计入第7项",
        "reason": "摘要存在报废可能且损益列入营业外支出",
        "supporting_evidence_ids": request["evidence_ids"],
        "contrary_evidence": [],
        "rejected_alternatives": [
            {"item_no": other_item, "reason": "没有正常出售或转让的明确证据"}
        ],
        "applicable_rule_ids": [
            rule["rule_id"]
            for rule in request["candidate_rules"]
            if rule["item_no"] == 7
        ],
        "cfo_impact_minor": request["available_amount_minor"],
        "confidence": 0.88,
        "strong_conflict": False,
        "human_review_required": False,
        "model_id": "cashflow-indirect-ai",
        "prompt_template_version": "cashflow-ai-v1",
        "prompt_hash": "a" * 64,
        "response_hash": "b" * 64,
    }
    path = tmp_path / "ai结论.jsonl"
    path.write_text(
        json.dumps(decision, ensure_ascii=False) + "\n",
        encoding="utf-8-sig",
    )
    return path


def test_prepare_stops_at_awaiting_ai_when_requests_exist(minimal_config):
    result = prepare_run(minimal_config)
    journal = pd.read_parquet(
        Path(result["run_dir"])
        / "02_normalized"
        / "journal_lines.parquet"
    )

    assert result["status"] == "awaiting_ai"
    assert result["ai_request_count"] == 1
    assert Path(result["ai_requests_path"]).exists()
    assert get_run_status(Path(result["run_dir"]))["status"] == "awaiting_ai"
    assert journal["source_account_code"].tolist() == ["671101", "160601"]


def test_prepare_does_not_send_ordinary_profit_loss_lines_to_ai(minimal_config):
    config = json.loads(Path(minimal_config).read_text(encoding="utf-8-sig"))
    workbook = openpyxl.load_workbook(config["journal_path"])
    journal = workbook[config["journal_sheet"]]
    journal.cell(row=2, column=6, value="支付办公费")
    journal.cell(row=2, column=7, value="6602")
    journal.cell(row=2, column=8, value="管理费用")
    journal.cell(row=2, column=9, value="660201")
    journal.cell(row=2, column=10, value="办公费")
    journal.cell(row=2, column=13, value="银行存款")
    journal.cell(row=3, column=6, value="支付办公费")
    journal.cell(row=3, column=7, value="1002")
    journal.cell(row=3, column=8, value="银行存款")
    journal.cell(row=3, column=9, value="100201")
    journal.cell(row=3, column=10, value="银行存款")
    journal.cell(row=3, column=13, value="办公费")

    balance = workbook[config["trial_balance_sheet"]]
    balance.cell(row=2, column=2, value="660201")
    balance.cell(row=2, column=3, value="办公费")
    balance.cell(row=2, column=6, value="损益")
    balance.cell(row=3, column=2, value="100201")
    balance.cell(row=3, column=3, value="银行存款")
    balance.cell(row=3, column=6, value="资产")
    workbook.save(config["journal_path"])
    workbook.close()

    result = prepare_run(minimal_config)

    assert result["status"] == "ready_to_finalize"
    assert result["ai_request_count"] == 0


def test_prepare_does_not_duplicate_deferred_tax_profit_loss_side(minimal_config):
    config = json.loads(Path(minimal_config).read_text(encoding="utf-8-sig"))
    workbook = openpyxl.load_workbook(config["journal_path"])
    journal = workbook[config["journal_sheet"]]
    for row in (2, 3):
        journal.cell(row=row, column=6, value="计提应收账款坏账准备")
        journal.cell(row=row, column=14, value="标准模板")
    journal.cell(row=2, column=7, value="1811")
    journal.cell(row=2, column=8, value="递延所得税资产")
    journal.cell(row=2, column=9, value="181101")
    journal.cell(row=2, column=10, value="递延所得税资产_应收账款坏账准备")
    journal.cell(row=2, column=13, value="所得税费用")
    journal.cell(row=3, column=7, value="6801")
    journal.cell(row=3, column=8, value="所得税费用")
    journal.cell(row=3, column=9, value="680102")
    journal.cell(row=3, column=10, value="所得税费用_递延所得税费用")
    journal.cell(row=3, column=13, value="递延所得税资产")

    balance = workbook[config["trial_balance_sheet"]]
    balance.cell(row=2, column=2, value="181101")
    balance.cell(row=2, column=3, value="递延所得税资产_应收账款坏账准备")
    balance.cell(row=2, column=6, value="资产")
    balance.cell(row=2, column=11, value=100)
    balance.cell(row=3, column=2, value="680102")
    balance.cell(row=3, column=3, value="所得税费用_递延所得税费用")
    balance.cell(row=3, column=6, value="损益")
    workbook.save(config["journal_path"])
    workbook.close()

    result = prepare_run(minimal_config)
    claims = pd.read_parquet(
        Path(result["run_dir"]) / "04_candidates" / "rule_claims.parquet"
    )

    assert result["status"] == "ready_to_finalize"
    assert result["ai_request_count"] == 0
    assert claims.loc[claims["item_no"].eq(11), "claim_amount_minor"].tolist() == [
        -10000
    ]
    assert claims.loc[claims["item_no"].eq(14)].empty


def test_prepare_direct_bad_debt_adjusts_item_two_and_receivable_once(
    minimal_config,
):
    config = json.loads(minimal_config.read_text(encoding="utf-8-sig"))
    workbook = openpyxl.load_workbook(config["journal_path"])
    journal = workbook[config["journal_sheet"]]
    direct_bad_debt_rows = [
        [
            "2026-12-31", "12月", "记", "001", "1", "直接核销坏账损失",
            "6602", "管理费用", "660201", "坏账损失", 100, 0,
            "应收账款", "算法生成(兜底拆分)",
        ],
        [
            "2026-12-31", "12月", "记", "001", "2", "直接核销坏账损失",
            "1122", "应收账款", "112201", "应收账款", 0, 100,
            "坏账损失", "算法生成(兜底拆分)",
        ],
    ]
    for row_index, values in enumerate(direct_bad_debt_rows, start=2):
        for column_index, value in enumerate(values, start=1):
            journal.cell(row=row_index, column=column_index, value=value)
    balance = workbook[config["trial_balance_sheet"]]
    balance_rows = [
        ["科目", "660201", "坏账损失", 1, "", "损益", 0, 0, 100, 0, 0, 0],
        ["科目", "112201", "应收账款", 1, "", "资产", 100, 0, 0, 100, 0, 0],
    ]
    for row_index, values in enumerate(balance_rows, start=2):
        for column_index, value in enumerate(values, start=1):
            balance.cell(row=row_index, column=column_index, value=value)
    workbook.save(config["journal_path"])
    workbook.close()

    prepared = prepare_run(minimal_config)
    claims = pd.read_parquet(
        Path(prepared["run_dir"]) / "04_candidates" / "rule_claims.parquet"
    )
    bridges = pd.read_parquet(
        Path(prepared["run_dir"]) / "04_candidates" / "family_bridges.parquet"
    )

    assert prepared["ai_request_count"] == 0
    assert claims.loc[claims["item_no"].eq(2), "claim_amount_minor"].sum() == 10000
    assert claims.loc[claims["item_no"].eq(14)].empty
    assert bridges.query(
        "family == 'working_capital' and item_no == 14 and record_type == 'summary'"
    )["final_amount_minor"].item() == 0


def test_prepare_deferred_tax_uses_counterparty_to_exclude_oci(
    minimal_config,
):
    config = json.loads(minimal_config.read_text(encoding="utf-8-sig"))
    workbook = openpyxl.load_workbook(config["journal_path"])
    journal = workbook[config["journal_sheet"]]
    journal_rows = [
        [
            "2026-12-31", "12月", "记", "001", "1", "期末确认税项",
            "1811", "递延所得税资产", "181101", "递延所得税资产", 100, 0,
            "其他综合收益", "算法生成(兜底拆分)",
        ],
        [
            "2026-12-31", "12月", "记", "001", "2", "期末确认税项",
            "4001", "其他综合收益", "400101", "其他综合收益", 0, 100,
            "递延所得税资产", "算法生成(兜底拆分)",
        ],
    ]
    for row_index, values in enumerate(journal_rows, start=2):
        for column_index, value in enumerate(values, start=1):
            journal.cell(row=row_index, column=column_index, value=value)
    balance = workbook[config["trial_balance_sheet"]]
    balance_rows = [
        ["科目", "181101", "递延所得税资产", 1, "", "资产", 0, 0, 100, 0, 100, 0],
        ["科目", "400101", "其他综合收益", 1, "", "权益", 0, 0, 0, 100, 0, 100],
    ]
    for row_index, values in enumerate(balance_rows, start=2):
        for column_index, value in enumerate(values, start=1):
            balance.cell(row=row_index, column=column_index, value=value)
    workbook.save(config["journal_path"])
    workbook.close()

    prepared = prepare_run(minimal_config)
    claims = pd.read_parquet(
        Path(prepared["run_dir"]) / "04_candidates" / "rule_claims.parquet"
    )
    bridges = pd.read_parquet(
        Path(prepared["run_dir"]) / "04_candidates" / "family_bridges.parquet"
    )

    assert prepared["ai_request_count"] == 0
    assert claims.loc[claims["item_no"].isin([11, 12])].empty
    assert bridges.query("family == 'deferred_tax'")["reason_code"].tolist() == [
        "NON_PROFIT_ROUTE"
    ]


def test_prepare_financial_entity_core_result_is_not_forced_into_general_rule(
    minimal_config,
):
    config = _update_config(minimal_config, financial_entity_flag=True)
    workbook = openpyxl.load_workbook(config["journal_path"])
    journal = workbook[config["journal_sheet"]]
    journal_rows = [
        [
            "2026-12-31", "12月", "记", "001", "1", "融资业务利息支出",
            "6603", "财务费用", "660301", "财务费用-利息支出", 100, 0,
            "银行存款", "算法生成(兜底拆分)",
        ],
        [
            "2026-12-31", "12月", "记", "001", "2", "融资业务利息支出",
            "1002", "银行存款", "100201", "银行存款", 0, 100,
            "财务费用-利息支出", "算法生成(兜底拆分)",
        ],
    ]
    for row_index, values in enumerate(journal_rows, start=2):
        for column_index, value in enumerate(values, start=1):
            journal.cell(row=row_index, column=column_index, value=value)
    balance = workbook[config["trial_balance_sheet"]]
    balance_rows = [
        ["科目", "660301", "财务费用-利息支出", 1, "", "损益", 0, 0, 100, 0, 0, 0],
        ["科目", "100201", "银行存款", 1, "", "资产", 100, 0, 0, 100, 0, 0],
    ]
    for row_index, values in enumerate(balance_rows, start=2):
        for column_index, value in enumerate(values, start=1):
            balance.cell(row=row_index, column=column_index, value=value)
    workbook.save(config["journal_path"])
    workbook.close()

    prepared = prepare_run(minimal_config)
    claims = pd.read_parquet(
        Path(prepared["run_dir"]) / "04_candidates" / "rule_claims.parquet"
    )

    assert prepared["status"] == "awaiting_ai"
    assert prepared["ai_request_count"] == 1
    assert claims.loc[claims["item_no"].isin([8, 9, 10])].empty


def test_prepare_share_payment_inventory_cost_reaches_item_thirteen_correction(
    minimal_config,
):
    config = json.loads(minimal_config.read_text(encoding="utf-8-sig"))
    workbook = openpyxl.load_workbook(config["journal_path"])
    journal = workbook[config["journal_sheet"]]
    journal_rows = [
        [
            "2026-12-31", "12月", "记", "001", "1",
            "权益结算股份支付形成未售存货", "1405", "库存商品", "140501",
            "库存商品", 60, 0, "资本公积-权益结算股份支付",
            "算法生成(兜底拆分)",
        ],
        [
            "2026-12-31", "12月", "记", "001", "2",
            "权益结算股份支付形成未售存货", "4002", "资本公积", "400201",
            "资本公积-权益结算股份支付", 0, 60, "库存商品",
            "算法生成(兜底拆分)",
        ],
    ]
    for row_index, values in enumerate(journal_rows, start=2):
        for column_index, value in enumerate(values, start=1):
            journal.cell(row=row_index, column=column_index, value=value)
    balance = workbook[config["trial_balance_sheet"]]
    balance_rows = [
        ["科目", "140501", "库存商品", 1, "", "资产", 0, 0, 60, 0, 60, 0],
        ["科目", "400201", "资本公积-权益结算股份支付", 1, "", "权益", 0, 0, 0, 60, 0, 60],
    ]
    for row_index, values in enumerate(balance_rows, start=2):
        for column_index, value in enumerate(values, start=1):
            balance.cell(row=row_index, column=column_index, value=value)
    workbook.save(config["journal_path"])
    workbook.close()

    prepared = prepare_run(minimal_config)
    claims = pd.read_parquet(
        Path(prepared["run_dir"]) / "04_candidates" / "rule_claims.parquet"
    )
    bridges = pd.read_parquet(
        Path(prepared["run_dir"]) / "04_candidates" / "family_bridges.parquet"
    )

    assert prepared["ai_request_count"] == 0
    assert claims.loc[claims["item_no"].isin([13, 16])].empty
    assert bridges.query(
        "family == 'other_adjustments' and target_item_no == 13"
    )["adjustment_minor"].item() == 6000
    assert bridges.query(
        "family == 'working_capital' and item_no == 13 and record_type == 'summary'"
    )["final_amount_minor"].item() == 0


def test_prepare_long_prepaid_one_time_non_cash_writeoff_reaches_item_five(
    minimal_config,
):
    config = json.loads(minimal_config.read_text(encoding="utf-8-sig"))
    workbook = openpyxl.load_workbook(config["journal_path"])
    journal = workbook[config["journal_sheet"]]
    journal_rows = [
        [
            "2026-12-31", "12月", "记", "001", "1",
            "长期待摊费用一次性非现金转销", "6602", "管理费用", "660201",
            "管理费用-长期待摊费用转销", 80, 0, "长期待摊费用",
            "算法生成(兜底拆分)",
        ],
        [
            "2026-12-31", "12月", "记", "001", "2",
            "长期待摊费用一次性非现金转销", "1801", "长期待摊费用", "180101",
            "长期待摊费用", 0, 80, "管理费用-长期待摊费用转销",
            "算法生成(兜底拆分)",
        ],
    ]
    for row_index, values in enumerate(journal_rows, start=2):
        for column_index, value in enumerate(values, start=1):
            journal.cell(row=row_index, column=column_index, value=value)
    balance = workbook[config["trial_balance_sheet"]]
    balance_rows = [
        ["科目", "660201", "管理费用-长期待摊费用转销", 1, "", "损益", 0, 0, 80, 0, 0, 0],
        ["科目", "180101", "长期待摊费用", 1, "", "资产", 80, 0, 0, 80, 0, 0],
    ]
    for row_index, values in enumerate(balance_rows, start=2):
        for column_index, value in enumerate(values, start=1):
            balance.cell(row=row_index, column=column_index, value=value)
    workbook.save(config["journal_path"])
    workbook.close()

    prepared = prepare_run(minimal_config)
    claims = pd.read_parquet(
        Path(prepared["run_dir"]) / "04_candidates" / "rule_claims.parquet"
    )

    assert prepared["ai_request_count"] == 0
    assert claims.query("item_no == 5")["claim_amount_minor"].item() == 8000
    assert claims.loc[claims["item_no"].isin([13, 14, 15])].empty


def test_finalize_consumes_valid_ai_decision_and_exports(minimal_config, tmp_path):
    prepared = prepare_run(minimal_config)
    ai_path = valid_ai_file(prepared, tmp_path)

    result = finalize_run(Path(prepared["run_dir"]), ai_path)

    assert result["status"] == "exported"
    assert result["completion_status"] == "formal_completed"
    assert result["unclassified_amount_minor"] == 0
    assert Path(result["workbook_path"]).exists()
    run_status = get_run_status(Path(prepared["run_dir"]))
    assert run_status["status"] == "exported"
    assert run_status["completion_status"] == "formal_completed"


def test_finalize_reports_real_journal_trial_balance_difference(
    minimal_config,
    tmp_path,
):
    config = json.loads(Path(minimal_config).read_text(encoding="utf-8-sig"))
    workbook = openpyxl.load_workbook(config["trial_balance_path"])
    worksheet = workbook[config["trial_balance_sheet"]]
    worksheet.cell(row=2, column=9, value=0)
    workbook.save(config["trial_balance_path"])
    workbook.close()

    prepared = prepare_run(minimal_config)
    ai_path = valid_ai_file(prepared, tmp_path)
    result = finalize_run(Path(prepared["run_dir"]), ai_path)

    reconciliation = json.loads(
        (
            Path(prepared["run_dir"])
            / "08_output"
            / "reconciliation.json"
        ).read_text(encoding="utf-8-sig")
    )
    check = next(
        item
        for item in reconciliation["checks"]
        if item["check_id"] == "journal_trial_balance"
    )
    assert check["difference_minor"] == 10000
    assert check["status"] == "fail"
    assert result["status"] == "blocked"
    assert result["file_status"] == "exported"
    assert result["completion_status"] == "diagnostic_only"
    assert Path(result["workbook_path"]).exists()
    result_summary = json.loads(
        (
            Path(prepared["run_dir"])
            / "08_output"
            / "result_summary.json"
        ).read_text(encoding="utf-8-sig")
    )
    assert result_summary["file_status"] == "exported"
    assert get_run_status(Path(prepared["run_dir"]))["status"] == "blocked"


def test_distinct_same_name_inputs_keep_separate_snapshots(
    minimal_config,
    tmp_path,
):
    config = json.loads(Path(minimal_config).read_text(encoding="utf-8-sig"))
    source = Path(config["journal_path"])
    journal_path = tmp_path / "序时账来源" / source.name
    balance_path = tmp_path / "余额表来源" / source.name
    journal_path.parent.mkdir()
    balance_path.parent.mkdir()
    shutil.copy2(source, journal_path)
    shutil.copy2(source, balance_path)

    workbook = openpyxl.load_workbook(balance_path)
    worksheet = workbook["序时账"]
    worksheet.cell(row=2, column=11, value=900)
    worksheet.cell(row=3, column=12, value=900)
    workbook.save(balance_path)
    workbook.close()

    config["journal_path"] = str(journal_path)
    config["trial_balance_path"] = str(balance_path)
    Path(minimal_config).write_text(
        json.dumps(config, ensure_ascii=False, indent=2),
        encoding="utf-8-sig",
    )
    prepared = prepare_run(minimal_config)
    journal = pd.read_parquet(
        Path(prepared["run_dir"])
        / "02_normalized"
        / "journal_lines.parquet"
    )

    assert int(journal["debit_minor"].sum()) == 10000
    snapshots = list(
        (Path(prepared["run_dir"]) / "01_input_snapshot").glob("*.xlsx")
    )
    assert len(snapshots) == 2


def test_rule_change_invalidates_prepared_checkpoint(minimal_config, tmp_path):
    config = json.loads(Path(minimal_config).read_text(encoding="utf-8-sig"))
    rule_source = Path(config["rules_path"])
    rule_copy = tmp_path / "本次规则.json"
    shutil.copy2(rule_source, rule_copy)
    config["rules_path"] = str(rule_copy)
    Path(minimal_config).write_text(
        json.dumps(config, ensure_ascii=False, indent=2),
        encoding="utf-8-sig",
    )
    prepared = prepare_run(minimal_config)

    rule_copy.write_text(
        rule_copy.read_text(encoding="utf-8-sig") + "\n",
        encoding="utf-8-sig",
    )
    status = get_run_status(Path(prepared["run_dir"]))

    assert status["status"] == "invalidated"
    assert status["invalidated_inputs"] == ["rules"]


def test_intermediate_file_change_invalidates_prepared_checkpoint(minimal_config):
    prepared = prepare_run(minimal_config)
    run_dir = Path(prepared["run_dir"])
    candidate_scope = run_dir / "04_candidates" / "candidate_scope.parquet"
    candidate_scope.write_bytes(candidate_scope.read_bytes() + b"changed")

    status = get_run_status(run_dir)

    assert status["status"] == "invalidated"
    assert "checkpoint" in status["invalidated_inputs"]


def test_ai_exclusion_finishes_without_false_unclassified_amount(
    minimal_config,
    tmp_path,
):
    _update_config(minimal_config, direct_cfo_minor=-10000)
    prepared = prepare_run(minimal_config)
    ai_path = valid_ai_file(prepared, tmp_path)
    decision = json.loads(ai_path.read_text(encoding="utf-8-sig"))
    decision.update(
        recommended_status="excluded",
        recommended_amount_minor=0,
        cfo_impact_minor=0,
        recommendation="排除该金额片段",
        reason="现有证据不足以证明属于非现金调节项目，按明确排除处理",
    )
    ai_path.write_text(
        json.dumps(decision, ensure_ascii=False) + "\n",
        encoding="utf-8-sig",
    )

    result = finalize_run(Path(prepared["run_dir"]), ai_path)

    assert result["status"] == "exported"
    assert result["completion_status"] == "formal_completed"
    assert result["unclassified_amount_minor"] == 0
    trace = pd.read_parquet(
        Path(prepared["run_dir"]) / "06_decisions" / "decision_trace.parquet"
    )
    assert "excluded" in set(trace["final_decision"])


def test_human_review_outputs_provisional_result_not_formal_completion(
    minimal_config,
    tmp_path,
):
    _update_config(minimal_config, performance_materiality_minor=10000)
    prepared = prepare_run(minimal_config)
    ai_path = valid_ai_file(prepared, tmp_path)
    decision = json.loads(ai_path.read_text(encoding="utf-8-sig"))
    decision.update(
        strong_conflict=True,
        human_review_required=True,
        contrary_evidence=["报废审批单与资产处置合同对业务性质的描述冲突"],
    )
    ai_path.write_text(
        json.dumps(decision, ensure_ascii=False) + "\n",
        encoding="utf-8-sig",
    )

    result = finalize_run(Path(prepared["run_dir"]), ai_path)

    assert result["status"] == "provisional_review"
    assert result["file_status"] == "exported"
    assert result["completion_status"] == "diagnostic_only"
    summary = pd.read_parquet(
        Path(prepared["run_dir"]) / "08_output" / "adjustment_summary.parquet"
    )
    item7 = summary.loc[summary["item_no"].eq(7)].iloc[0]
    assert item7["issue_count"] == 1
    assert item7["check_status"] == "review_required"
    validated = Path(prepared["run_dir"]) / "05_ai" / "ai_decisions_validated.jsonl"
    assert validated.is_file()


def test_run_state_records_frozen_journal_for_workpaper(minimal_config):
    prepared = prepare_run(minimal_config)
    run_dir = Path(prepared["run_dir"])
    state = json.loads(
        (run_dir / "run_state.json").read_text("utf-8-sig")
    )
    snapshot = Path(state["snapshot_sources"]["journal"])
    assert snapshot.is_file()
    assert snapshot.parent == run_dir / "01_input_snapshot"
    assert state["journal_sheet"] == "序时账"
    assert state["journal_header_row"] == 1
    assert state["journal_row_count"] > 0
    assert state["journal_column_count"] > 0


def test_snapshot_missing_or_changed_invalidates_with_clear_error(
    minimal_config,
):
    prepared = prepare_run(minimal_config)
    run_dir = Path(prepared["run_dir"])
    state = json.loads(
        (run_dir / "run_state.json").read_text("utf-8-sig")
    )
    snapshot = Path(state["snapshot_sources"]["journal"])
    assert snapshot.is_file()

    # 场景一：快照被修改
    snapshot.write_bytes(snapshot.read_bytes() + b"tampered")
    status = get_run_status(run_dir)
    assert status["status"] == "invalidated"
    assert "journal" in status["invalidated_inputs"]
    assert "冻结输入快照缺失或已变化" in status["invalidation_reason"]

    # 场景二：快照被删除
    snapshot.unlink()
    status = get_run_status(run_dir)
    assert status["status"] == "invalidated"
    assert "journal" in status["invalidated_inputs"]
    assert "冻结输入快照缺失或已变化" in status["invalidation_reason"]



def test_finalize_saves_decision_trace_with_reason_and_occupancy(
    minimal_config,
    tmp_path,
):
    prepared = prepare_run(minimal_config)
    ai_path = valid_ai_file(prepared, tmp_path)

    result = finalize_run(Path(prepared["run_dir"]), ai_path)

    assert result["status"] == "exported"
    trace = pd.read_parquet(
        Path(prepared["run_dir"]) / "06_decisions" / "decision_trace.parquet"
    )
    assert not trace.empty
    assert {"trace_id", "fragment_id", "item_no", "final_decision"}.issubset(
        set(trace.columns)
    )
    assert trace["decision_reason"].fillna("").astype(str).str.len().gt(0).all()
    assert trace["physical_consumed_minor"].sum() > 0


def test_prepare_rejects_invalid_output_mode(minimal_config):
    _update_config(minimal_config, output_mode="turbo")

    with pytest.raises(InputProtocolError, match="输出模式"):
        prepare_run(minimal_config)


@pytest.mark.parametrize(
    ("field", "value", "message"),
    [
        ("entity_name", "", "主体名称"),
        ("functional_currency", "", "本位币"),
        ("currency_scale", -1, "小数位数"),
        ("display_unit", "亿元", "展示单位"),
        ("performance_materiality_minor", -1, "重要性水平"),
        ("financial_entity_flag", "false", "金融企业标志"),
    ],
)
def test_prepare_rejects_invalid_run_spec(
    minimal_config,
    field,
    value,
    message,
):
    _update_config(minimal_config, **{field: value})

    with pytest.raises(InputProtocolError, match=message):
        prepare_run(minimal_config)


def test_prepare_rejects_missing_financial_entity_flag(minimal_config):
    config = json.loads(minimal_config.read_text(encoding="utf-8-sig"))
    config.pop("financial_entity_flag", None)
    minimal_config.write_text(
        json.dumps(config, ensure_ascii=False, indent=2),
        encoding="utf-8-sig",
    )

    with pytest.raises(InputProtocolError, match="金融企业标志"):
        prepare_run(minimal_config)


def test_prepare_rejects_legacy_excel_format_before_reading(
    minimal_config,
    tmp_path,
):
    legacy = tmp_path / "旧版序时账.xls"
    legacy.write_bytes(b"legacy")
    _update_config(minimal_config, journal_path=str(legacy))

    with pytest.raises(InputProtocolError, match="xlsx或xlsm"):
        prepare_run(minimal_config)


def test_finalize_uses_rapid_output_mode_from_config(minimal_config, tmp_path):
    _update_config(minimal_config, output_mode="rapid")
    prepared = prepare_run(minimal_config)
    ai_path = valid_ai_file(prepared, tmp_path)

    result = finalize_run(Path(prepared["run_dir"]), ai_path)

    assert result["status"] == "exported"
    run_dir = Path(prepared["run_dir"])
    state = json.loads(
        (run_dir / "run_state.json").read_text(encoding="utf-8-sig")
    )
    assert state["output_mode"] == "rapid"

    workbook = openpyxl.load_workbook(result["workbook_path"], read_only=True)
    assert workbook.sheetnames == ["附表结果", "简报", "机器用复核数据"]
    brief = workbook["简报"]
    brief_map = {
        row[0].value: row[1].value for row in brief.iter_rows(max_col=2)
    }
    workbook.close()
    assert brief_map["运行标识"] == state["run_id"]
    assert brief_map["规则版本"] == state["rule_version"]
    assert brief_map["代码版本"] == state["code_version"]

    md_dir = run_dir / "08_output" / "底稿MD"
    md_files = sorted(md_dir.glob("*.md"))
    assert len(md_files) == 16
    manifest = json.loads(
        (run_dir / "08_output" / "output_manifest.json").read_text(
            encoding="utf-8-sig"
        )
    )
    manifest_paths = {entry["path"] for entry in manifest["files"]}
    assert all(f"底稿MD/{path.name}" in manifest_paths for path in md_files)


def test_finalize_defaults_to_standard_output_mode(minimal_config, tmp_path):
    prepared = prepare_run(minimal_config)
    ai_path = valid_ai_file(prepared, tmp_path)

    result = finalize_run(Path(prepared["run_dir"]), ai_path)

    assert result["status"] == "exported"
    workbook = openpyxl.load_workbook(result["workbook_path"], read_only=True)
    sheet_count = len(workbook.sheetnames)
    workbook.close()
    assert sheet_count == 19
    assert not (Path(prepared["run_dir"]) / "08_output" / "底稿MD").exists()


def test_finalize_without_direct_cfo_exports_diagnostic_only(
    minimal_config,
    tmp_path,
):
    _update_config(minimal_config, direct_cfo_minor=None)
    prepared = prepare_run(minimal_config)
    ai_path = valid_ai_file(prepared, tmp_path)

    result = finalize_run(Path(prepared["run_dir"]), ai_path)

    assert result["status"] == "exported"
    assert result["file_status"] == "exported"
    assert result["completion_status"] == "diagnostic_only"
    assert Path(result["workbook_path"]).exists()
    run_state = get_run_status(Path(prepared["run_dir"]))
    assert run_state["completion_status"] == "diagnostic_only"
    result_summary = json.loads(
        (
            Path(prepared["run_dir"])
            / "08_output"
            / "result_summary.json"
        ).read_text(encoding="utf-8-sig")
    )
    assert result_summary["file_status"] == "exported"
    assert result_summary["completion_status"] == "diagnostic_only"
    reconciliation = json.loads(
        (
            Path(prepared["run_dir"])
            / "08_output"
            / "reconciliation.json"
        ).read_text(encoding="utf-8-sig")
    )
    assert reconciliation["status"] == "not_run"
    workbook = openpyxl.load_workbook(result["workbook_path"], read_only=True)
    try:
        result_rows = {
            row[0].value: row[1].value
            for row in workbook["附表结果"].iter_rows(min_col=1, max_col=2)
            if row[0].value
        }
    finally:
        workbook.close()
    assert result_rows["成果性质"] == "仅诊断、待总对账"


def test_blocked_run_still_writes_diagnostic_files(minimal_config, tmp_path):
    prepared = prepare_run(minimal_config)
    request_path = Path(prepared["ai_requests_path"])
    request = json.loads(request_path.read_text("utf-8-sig").splitlines()[0])
    other_item = next(
        item for item in request["candidate_items"] if item != 7
    )
    decision = {
        "request_id": request["request_id"],
        "recommended_item": 7,
        "recommended_status": "included",
        "recommended_amount_minor": request["available_amount_minor"] - 1,
        "recommendation": "作为固定资产报废损失计入第7项",
        "reason": "摘要存在报废可能且损益列入营业外支出",
        "supporting_evidence_ids": request["evidence_ids"],
        "contrary_evidence": [],
        "rejected_alternatives": [
            {"item_no": other_item, "reason": "没有正常出售或转让的明确证据"}
        ],
        "applicable_rule_ids": [
            rule["rule_id"]
            for rule in request["candidate_rules"]
            if rule["item_no"] == 7
        ],
        "cfo_impact_minor": request["available_amount_minor"] - 1,
        "confidence": 0.88,
        "strong_conflict": False,
        "human_review_required": False,
        "model_id": "cashflow-indirect-ai",
        "prompt_template_version": "cashflow-ai-v1",
        "prompt_hash": "a" * 64,
        "response_hash": "b" * 64,
    }
    ai_path = tmp_path / "ai结论少一分.jsonl"
    ai_path.write_text(
        json.dumps(decision, ensure_ascii=False) + "\n",
        encoding="utf-8-sig",
    )

    result = finalize_run(Path(prepared["run_dir"]), ai_path)

    assert result["status"] == "blocked"
    assert result["unclassified_amount_minor"] == 1
    output_dir = Path(prepared["run_dir"]) / "08_output"
    diagnosis = json.loads(
        (output_dir / "reconciliation.json").read_text(encoding="utf-8-sig")
    )
    assert diagnosis["status"] == "blocked"
    assert diagnosis["unclassified_amount_minor"] == 1
    assert diagnosis["next_action"]
    issues_text = (output_dir / "issues.jsonl").read_text(encoding="utf-8-sig")
    assert "UNCLASSIFIED_REMAINDER" in issues_text


def test_share_payment_inventory_finalize_balance_bridge_does_not_double_count(
    minimal_config,
):
    config = json.loads(minimal_config.read_text(encoding="utf-8-sig"))
    workbook = openpyxl.load_workbook(config["journal_path"])
    journal = workbook[config["journal_sheet"]]
    journal_rows = [
        [
            "2026-12-31", "12月", "记", "001", "1",
            "权益结算股份支付形成未售存货", "1405", "库存商品", "140501",
            "库存商品", 60, 0, "资本公积-权益结算股份支付", "算法生成(兜底拆分)",
        ],
        [
            "2026-12-31", "12月", "记", "001", "2",
            "权益结算股份支付形成未售存货", "4002", "资本公积", "400201",
            "资本公积-权益结算股份支付", 0, 60, "库存商品", "算法生成(兜底拆分)",
        ],
    ]
    for row_index, values in enumerate(journal_rows, start=2):
        for column_index, value in enumerate(values, start=1):
            journal.cell(row=row_index, column=column_index, value=value)
    balance = workbook[config["trial_balance_sheet"]]
    balance_rows = [
        ["科目", "140501", "库存商品", 1, "", "资产", 0, 0, 60, 0, 60, 0],
        ["科目", "400201", "资本公积-权益结算股份支付", 1, "", "权益", 0, 0, 0, 60, 0, 60],
    ]
    for row_index, values in enumerate(balance_rows, start=2):
        for column_index, value in enumerate(values, start=1):
            balance.cell(row=row_index, column=column_index, value=value)
    workbook.save(config["journal_path"])
    workbook.close()
    _update_config(minimal_config, net_profit_minor=0, direct_cfo_minor=0)

    prepared = prepare_run(minimal_config)
    assert prepared["ai_request_count"] == 0
    result = finalize_run(Path(prepared["run_dir"]), None)

    assert result["status"] == "exported"
    reconciliation = json.loads(
        (Path(prepared["run_dir"]) / "08_output" / "reconciliation.json")
        .read_text(encoding="utf-8-sig")
    )
    balance_bridge = next(
        check for check in reconciliation["checks"]
        if check["check_id"] == "balance_bridge"
    )
    assert balance_bridge["status"] == "pass"
    assert balance_bridge["difference_minor"] == 0


def test_abnormal_intangible_amortization_issue_reaches_formal_entry(
    minimal_config,
):
    config = json.loads(minimal_config.read_text(encoding="utf-8-sig"))
    workbook = openpyxl.load_workbook(config["journal_path"])
    journal = workbook[config["journal_sheet"]]
    journal_rows = [
        [
            "2026-12-31", "12月", "记", "001", "1",
            "使用寿命不确定的无形资产错误摊销", "6602", "管理费用", "660201",
            "管理费用-无形资产摊销", 10, 0, "累计摊销", "算法生成(兜底拆分)",
        ],
        [
            "2026-12-31", "12月", "记", "001", "2",
            "使用寿命不确定的无形资产错误摊销", "1701", "无形资产", "170102",
            "累计摊销", 0, 10, "管理费用", "算法生成(兜底拆分)",
        ],
    ]
    for row_index, values in enumerate(journal_rows, start=2):
        for column_index, value in enumerate(values, start=1):
            journal.cell(row=row_index, column=column_index, value=value)
    balance = workbook[config["trial_balance_sheet"]]
    balance_rows = [
        ["科目", "660201", "管理费用-无形资产摊销", 1, "", "损益", 0, 0, 10, 0, 0, 0],
        ["科目", "170102", "累计摊销", 1, "", "资产", 0, 0, 0, 10, 0, 10],
    ]
    for row_index, values in enumerate(balance_rows, start=2):
        for column_index, value in enumerate(values, start=1):
            balance.cell(row=row_index, column=column_index, value=value)
    workbook.save(config["journal_path"])
    workbook.close()

    prepared = prepare_run(minimal_config)
    issues = pd.read_parquet(
        Path(prepared["run_dir"]) / "04_candidates" / "family_issues.parquet"
    )
    claims = pd.read_parquet(
        Path(prepared["run_dir"]) / "04_candidates" / "rule_claims.parquet"
    )

    assert (issues["reason_code"] == "INTANGIBLE_AMORTIZATION_NOT_ALLOWED").any()
    assert claims.loc[claims["item_no"].eq(4), "claim_amount_minor"].sum() == 1000


def test_debt_restructuring_profit_reaches_item_sixteen(minimal_config):
    config = json.loads(minimal_config.read_text(encoding="utf-8-sig"))
    workbook = openpyxl.load_workbook(config["journal_path"])
    journal = workbook[config["journal_sheet"]]
    journal_rows = [
        [
            "2026-12-31", "12月", "记", "001", "1",
            "债务重组以资抵债确认收益", "2201", "应付账款", "220101",
            "应付账款", 100, 0, "营业外收入", "算法生成(兜底拆分)",
        ],
        [
            "2026-12-31", "12月", "记", "001", "2",
            "债务重组以资抵债确认收益", "6301", "营业外收入", "630101",
            "营业外收入-债务重组利得", 0, 60, "应付账款", "算法生成(兜底拆分)",
        ],
        [
            "2026-12-31", "12月", "记", "001", "3",
            "债务重组以资抵债确认收益", "1002", "银行存款", "100201",
            "银行存款", 0, 40, "应付账款", "算法生成(兜底拆分)",
        ],
    ]
    for row_index, values in enumerate(journal_rows, start=2):
        for column_index, value in enumerate(values, start=1):
            journal.cell(row=row_index, column=column_index, value=value)
    balance = workbook[config["trial_balance_sheet"]]
    balance_rows = [
        ["科目", "220101", "应付账款", 1, "", "负债", 0, 100, 100, 0, 0, 0],
        ["科目", "630101", "营业外收入-债务重组利得", 1, "", "损益", 0, 0, 0, 60, 0, 0],
        ["科目", "100201", "银行存款", 1, "", "资产", 100, 0, 0, 40, 60, 0],
    ]
    for row_index, values in enumerate(balance_rows, start=2):
        for column_index, value in enumerate(values, start=1):
            balance.cell(row=row_index, column=column_index, value=value)
    workbook.save(config["journal_path"])
    workbook.close()

    prepared = prepare_run(minimal_config)
    claims = pd.read_parquet(
        Path(prepared["run_dir"]) / "04_candidates" / "rule_claims.parquet"
    )

    assert prepared["ai_request_count"] == 0
    assert claims.loc[claims["item_no"].eq(16), "claim_amount_minor"].sum() == -6000
    assert claims.loc[claims["item_no"].eq(15), "claim_amount_minor"].sum() == -4000


def test_debt_restructuring_inventory_settlement_corrects_working_capital(
    minimal_config,
):
    config = json.loads(minimal_config.read_text(encoding="utf-8-sig"))
    workbook = openpyxl.load_workbook(config["journal_path"])
    journal = workbook[config["journal_sheet"]]
    journal_rows = [
        [
            "2026-12-31", "12月", "记", "001", "1",
            "债务重组以存货抵债", "2201", "应付账款", "220101",
            "应付账款", 100, 0, "库存商品", "算法生成(兜底拆分)",
        ],
        [
            "2026-12-31", "12月", "记", "001", "2",
            "债务重组以存货抵债", "1405", "库存商品", "140501",
            "库存商品", 0, 80, "应付账款", "算法生成(兜底拆分)",
        ],
        [
            "2026-12-31", "12月", "记", "001", "3",
            "债务重组以存货抵债", "6301", "营业外收入", "630101",
            "营业外收入-债务重组利得", 0, 20, "应付账款", "算法生成(兜底拆分)",
        ],
    ]
    for row_index, values in enumerate(journal_rows, start=2):
        for column_index, value in enumerate(values, start=1):
            journal.cell(row=row_index, column=column_index, value=value)
    balance = workbook[config["trial_balance_sheet"]]
    balance_rows = [
        ["科目", "220101", "应付账款", 1, "", "负债", 0, 100, 100, 0, 0, 0],
        ["科目", "140501", "库存商品", 1, "", "资产", 80, 0, 0, 80, 0, 0],
        ["科目", "630101", "营业外收入-债务重组利得", 1, "", "损益", 0, 0, 0, 20, 0, 0],
    ]
    for row_index, values in enumerate(balance_rows, start=2):
        for column_index, value in enumerate(values, start=1):
            balance.cell(row=row_index, column=column_index, value=value)
    workbook.save(config["journal_path"])
    workbook.close()

    prepared = prepare_run(minimal_config)
    claims = pd.read_parquet(
        Path(prepared["run_dir"]) / "04_candidates" / "rule_claims.parquet"
    )

    assert prepared["ai_request_count"] == 0
    assert claims.loc[claims["item_no"].eq(13)].empty
    assert claims.loc[claims["item_no"].eq(15)].empty
    assert claims.loc[claims["item_no"].eq(16), "claim_amount_minor"].sum() == -2000


def test_shareholder_debt_waiver_is_not_ordinary_debt_restructuring(minimal_config):
    config = json.loads(minimal_config.read_text(encoding="utf-8-sig"))
    workbook = openpyxl.load_workbook(config["journal_path"])
    journal = workbook[config["journal_sheet"]]
    journal_rows = [
        [
            "2026-12-31", "12月", "记", "001", "1",
            "股东身份豁免债务作为权益性交易", "2201", "应付账款", "220101",
            "应付账款", 100, 0, "资本公积", "算法生成(兜底拆分)",
        ],
        [
            "2026-12-31", "12月", "记", "001", "2",
            "股东身份豁免债务作为权益性交易", "4002", "资本公积", "400201",
            "资本公积-股东豁免", 0, 100, "应付账款", "算法生成(兜底拆分)",
        ],
    ]
    for row_index, values in enumerate(journal_rows, start=2):
        for column_index, value in enumerate(values, start=1):
            journal.cell(row=row_index, column=column_index, value=value)
    balance = workbook[config["trial_balance_sheet"]]
    balance_rows = [
        ["科目", "220101", "应付账款", 1, "", "负债", 0, 100, 100, 0, 0, 0],
        ["科目", "400201", "资本公积-股东豁免", 1, "", "权益", 0, 0, 0, 100, 0, 100],
    ]
    for row_index, values in enumerate(balance_rows, start=2):
        for column_index, value in enumerate(values, start=1):
            balance.cell(row=row_index, column=column_index, value=value)
    workbook.save(config["journal_path"])
    workbook.close()

    prepared = prepare_run(minimal_config)
    claims = pd.read_parquet(
        Path(prepared["run_dir"]) / "04_candidates" / "rule_claims.parquet"
    )
    bridges = pd.read_parquet(
        Path(prepared["run_dir"]) / "04_candidates" / "family_bridges.parquet"
    )

    assert claims.loc[claims["item_no"].eq(16)].empty
    assert claims.loc[claims["item_no"].eq(15)].empty
    assert bridges.query(
        "family == 'other_adjustments' and reason_code == 'SHAREHOLDER_EQUITY_TRANSACTION'"
    ).shape[0] >= 1


def test_deferred_tax_balance_fallback_reaches_formal_entry(minimal_config):
    config = json.loads(minimal_config.read_text(encoding="utf-8-sig"))
    workbook = openpyxl.load_workbook(config["journal_path"])
    journal = workbook[config["journal_sheet"]]
    journal_rows = [
        [
            "2026-12-31", "12月", "记", "001", "1",
            "支付办公费用", "6602", "管理费用", "660299",
            "管理费用-办公费", 10, 0, "银行存款", "算法生成(兜底拆分)",
        ],
        [
            "2026-12-31", "12月", "记", "001", "2",
            "支付办公费用", "1002", "银行存款", "100201",
            "银行存款", 0, 10, "管理费用", "算法生成(兜底拆分)",
        ],
    ]
    for row_index, values in enumerate(journal_rows, start=2):
        for column_index, value in enumerate(values, start=1):
            journal.cell(row=row_index, column=column_index, value=value)
    balance = workbook[config["trial_balance_sheet"]]
    balance_rows = [
        ["科目", "660299", "管理费用-办公费", 1, "", "损益", 0, 0, 10, 0, 0, 0],
        ["科目", "100201", "银行存款", 1, "", "资产", 0, 0, 0, 10, 0, 0],
        ["科目", "181101", "递延所得税资产", 1, "", "资产", 1000, 0, 0, 0, 400, 0],
    ]
    for row_index, values in enumerate(balance_rows, start=2):
        for column_index, value in enumerate(values, start=1):
            balance.cell(row=row_index, column=column_index, value=value)
    workbook.save(config["journal_path"])
    workbook.close()

    prepared = prepare_run(minimal_config)
    claims = pd.read_parquet(
        Path(prepared["run_dir"]) / "04_candidates" / "rule_claims.parquet"
    )

    assert prepared["ai_request_count"] == 0
    assert claims.loc[claims["item_no"].eq(11), "claim_amount_minor"].sum() == 60000


def test_ordinary_bank_deposit_interest_is_not_added_to_item_nine(minimal_config):
    config = json.loads(minimal_config.read_text(encoding="utf-8-sig"))
    workbook = openpyxl.load_workbook(config["journal_path"])
    journal = workbook[config["journal_sheet"]]
    journal_rows = [
        [
            "2026-12-31", "12月", "记", "001", "1",
            "银行存款利息收入", "1002", "银行存款", "100201",
            "银行存款", 8, 0, "财务费用", "算法生成(兜底拆分)",
        ],
        [
            "2026-12-31", "12月", "记", "001", "2",
            "银行存款利息收入", "6603", "财务费用", "660301",
            "财务费用-存款利息收入", 0, 8, "银行存款", "算法生成(兜底拆分)",
        ],
    ]
    for row_index, values in enumerate(journal_rows, start=2):
        for column_index, value in enumerate(values, start=1):
            journal.cell(row=row_index, column=column_index, value=value)
    balance = workbook[config["trial_balance_sheet"]]
    balance_rows = [
        ["科目", "100201", "银行存款", 1, "", "资产", 0, 0, 8, 0, 8, 0],
        ["科目", "660301", "财务费用-存款利息收入", 1, "", "损益", 0, 0, 0, 8, 0, 0],
    ]
    for row_index, values in enumerate(balance_rows, start=2):
        for column_index, value in enumerate(values, start=1):
            balance.cell(row=row_index, column=column_index, value=value)
    workbook.save(config["journal_path"])
    workbook.close()

    prepared = prepare_run(minimal_config)
    claims = pd.read_parquet(
        Path(prepared["run_dir"]) / "04_candidates" / "rule_claims.parquet"
    )

    assert prepared["ai_request_count"] == 0
    assert claims.loc[claims["item_no"].eq(9)].empty


def test_investment_product_interest_recorded_in_finance_cost_goes_item_nine(
    minimal_config,
):
    config = json.loads(minimal_config.read_text(encoding="utf-8-sig"))
    workbook = openpyxl.load_workbook(config["journal_path"])
    journal = workbook[config["journal_sheet"]]
    journal_rows = [
        [
            "2026-12-31", "12月", "记", "001", "1",
            "理财产品利息收入", "1002", "银行存款", "100201",
            "银行存款", 8, 0, "财务费用", "算法生成(兜底拆分)",
        ],
        [
            "2026-12-31", "12月", "记", "001", "2",
            "理财产品利息收入", "6603", "财务费用", "660301",
            "财务费用-理财产品利息收入", 0, 8, "银行存款", "算法生成(兜底拆分)",
        ],
    ]
    for row_index, values in enumerate(journal_rows, start=2):
        for column_index, value in enumerate(values, start=1):
            journal.cell(row=row_index, column=column_index, value=value)
    balance = workbook[config["trial_balance_sheet"]]
    balance_rows = [
        ["科目", "100201", "银行存款", 1, "", "资产", 0, 0, 8, 0, 8, 0],
        ["科目", "660301", "财务费用-理财产品利息收入", 1, "", "损益", 0, 0, 0, 8, 0, 0],
    ]
    for row_index, values in enumerate(balance_rows, start=2):
        for column_index, value in enumerate(values, start=1):
            balance.cell(row=row_index, column=column_index, value=value)
    workbook.save(config["journal_path"])
    workbook.close()

    prepared = prepare_run(minimal_config)
    claims = pd.read_parquet(
        Path(prepared["run_dir"]) / "04_candidates" / "rule_claims.parquet"
    )

    assert prepared["ai_request_count"] == 0
    assert claims.loc[claims["item_no"].eq(9), "claim_amount_minor"].sum() == -800


def test_recovery_of_written_off_amount_is_not_treated_as_reversal(minimal_config):
    config = json.loads(minimal_config.read_text(encoding="utf-8-sig"))
    workbook = openpyxl.load_workbook(config["journal_path"])
    journal = workbook[config["journal_sheet"]]
    journal_rows = [
        [
            "2026-12-31", "12月", "记", "001", "1",
            "收回已核销坏账款项", "1002", "银行存款", "100201",
            "银行存款", 10, 0, "信用减值损失", "算法生成(兜底拆分)",
        ],
        [
            "2026-12-31", "12月", "记", "001", "2",
            "收回已核销坏账款项", "6702", "信用减值损失", "670201",
            "信用减值损失-收回已核销", 0, 10, "银行存款", "算法生成(兜底拆分)",
        ],
    ]
    for row_index, values in enumerate(journal_rows, start=2):
        for column_index, value in enumerate(values, start=1):
            journal.cell(row=row_index, column=column_index, value=value)
    balance = workbook[config["trial_balance_sheet"]]
    balance_rows = [
        ["科目", "100201", "银行存款", 1, "", "资产", 0, 0, 10, 0, 10, 0],
        ["科目", "670201", "信用减值损失-收回已核销", 1, "", "损益", 0, 0, 0, 10, 0, 0],
    ]
    for row_index, values in enumerate(balance_rows, start=2):
        for column_index, value in enumerate(values, start=1):
            balance.cell(row=row_index, column=column_index, value=value)
    workbook.save(config["journal_path"])
    workbook.close()

    prepared = prepare_run(minimal_config)
    claims = pd.read_parquet(
        Path(prepared["run_dir"]) / "04_candidates" / "rule_claims.parquet"
    )

    assert prepared["ai_request_count"] == 0
    assert claims.loc[claims["item_no"].eq(2)].empty


def test_debt_restructuring_does_not_double_claim_investment_result(minimal_config):
    config = json.loads(minimal_config.read_text(encoding="utf-8-sig"))
    workbook = openpyxl.load_workbook(config["journal_path"])
    journal = workbook[config["journal_sheet"]]
    journal_rows = [
        [
            "2026-12-31", "12月", "记", "001", "1",
            "债务重组以股权投资清偿", "2201", "应付账款", "220101",
            "应付账款", 100, 0, "长期股权投资", "算法生成(兜底拆分)",
        ],
        [
            "2026-12-31", "12月", "记", "001", "2",
            "债务重组以股权投资清偿", "1511", "长期股权投资", "151101",
            "长期股权投资", 0, 70, "应付账款", "算法生成(兜底拆分)",
        ],
        [
            "2026-12-31", "12月", "记", "001", "3",
            "债务重组以股权投资清偿", "6111", "投资收益", "611101",
            "投资收益-债务重组利得", 0, 30, "应付账款", "算法生成(兜底拆分)",
        ],
    ]
    for row_index, values in enumerate(journal_rows, start=2):
        for column_index, value in enumerate(values, start=1):
            journal.cell(row=row_index, column=column_index, value=value)
    balance = workbook[config["trial_balance_sheet"]]
    balance_rows = [
        ["科目", "220101", "应付账款", 1, "", "负债", 0, 100, 100, 0, 0, 0],
        ["科目", "151101", "长期股权投资", 1, "", "资产", 70, 0, 0, 70, 0, 0],
        ["科目", "611101", "投资收益-债务重组利得", 1, "", "损益", 0, 0, 0, 30, 0, 0],
    ]
    for row_index, values in enumerate(balance_rows, start=2):
        for column_index, value in enumerate(values, start=1):
            balance.cell(row=row_index, column=column_index, value=value)
    workbook.save(config["journal_path"])
    workbook.close()

    prepared = prepare_run(minimal_config)
    claims = pd.read_parquet(
        Path(prepared["run_dir"]) / "04_candidates" / "rule_claims.parquet"
    )

    assert prepared["ai_request_count"] == 0
    assert claims.loc[claims["item_no"].eq(10), "claim_amount_minor"].sum() == -3000
    assert claims.loc[claims["item_no"].eq(16)].empty
    assert claims.loc[claims["item_no"].eq(15)].empty
