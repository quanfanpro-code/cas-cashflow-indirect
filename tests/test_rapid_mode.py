"""急速模式输出测试：2 个 sheet 的 xlsx + 16 个同名 MD 底稿。"""

from types import SimpleNamespace

import openpyxl
import pandas as pd
import pytest

from cashflow_indirect.reconciliation import calculate_summary, reconcile_all
from cashflow_indirect.workpaper import write_item_markdowns, write_rapid_workbook


@pytest.fixture
def rapid_input(tmp_path, rule_book):
    snapshot = tmp_path / "序时账快照.xlsx"
    workbook = openpyxl.Workbook()
    worksheet = workbook.active
    worksheet.title = "序时账"
    worksheet.append(
        [
            "记账时间", "会计月", "凭证种类", "凭证编号", "编号",
            "业务说明", "科目编号", "科目名称", "借方发生额", "贷方发生额",
            "对方科目", "匹配类型",
        ]
    )
    worksheet.append(
        [
            "2026-12-31", "12月", "记", "001", "1",
            "计提资产减值准备", "671101", "资产减值损失",
            20000, 0, "存货跌价准备", "算法生成",
        ]
    )
    worksheet.append(
        [
            "2026-12-31", "12月", "记", "001", "2",
            "结转损益", "410401", "本年利润",
            0, 20000, "资产减值损失", "结转损益规则",
        ]
    )
    worksheet.append(
        [
            "2026-12-31", "12月", "记", "002", "3",
            "支付办公费", "660201", "办公费",
            5000, 0, "银行存款", "算法生成",
        ]
    )
    workbook.save(snapshot)
    workbook.close()

    occupancy = pd.DataFrame(
        [
            {
                "fragment_id": "f-1",
                "item_no": 1,
                "adjustment_amount_minor": 20000,
                "claimed_amount_minor": 20000,
                "physical_consumed_minor": 20000,
                "remaining_amount_minor": 0,
                "decision": "included",
                "rule_id": "IMP-ASSET",
                "paired_offset_id": None,
            }
        ]
    )
    summary = calculate_summary(1_000_000, occupancy, rule_book)
    reconciliation = reconcile_all(summary, direct_cfo_minor=1_020_000)
    details = pd.DataFrame(
        [
            {
                "item_no": 1,
                "item_name": "资产减值准备",
                "business_event_id": "event-1",
                "fragment_id": "f-1",
                "voucher_date": "2026-12-31",
                "voucher_no": "记-001",
                "source_row": 2,
                "account_name": "资产减值损失",
                "counterparty_level1": "存货跌价准备",
                "summary": "计提跌价准备",
                "source_amount_minor": 20000,
                "adjustment_amount_minor": 20000,
                "decision": "included",
                "rule_id": "IMP-ASSET",
                "reason": "计入当期损益的减值",
            }
        ]
    )
    bridges = pd.DataFrame()
    issues = pd.DataFrame()
    decision_trace = pd.DataFrame(
        [
            {
                "trace_id": "trace-1",
                "fragment_id": "f-1",
                "business_event_id": "event-1",
                "source_row": 2,
                "item_no": 1,
                "candidate_status": "rule_classified",
                "final_decision": "included",
                "rule_id": "IMP-ASSET",
                "decision_reason": "计入当期损益的减值",
                "confidence": 1.0,
                "recommendation": "",
                "recommendation_reason": "",
                "supporting_evidence": "",
                "contrary_evidence": "未发现相反证据",
                "rejected_alternatives": "",
                "human_review_required": False,
                "claim_amount_minor": 20000,
                "physical_consumed_minor": 20000,
                "remaining_amount_minor": 0,
                "occupancy_type": "exclusive",
                "paired_offset_id": None,
            }
        ],
        columns=[
            "trace_id", "fragment_id", "business_event_id", "source_row",
            "item_no", "candidate_status", "final_decision", "rule_id",
            "decision_reason", "confidence", "recommendation",
            "recommendation_reason", "supporting_evidence",
            "contrary_evidence", "rejected_alternatives",
            "human_review_required", "claim_amount_minor",
            "physical_consumed_minor", "remaining_amount_minor",
            "occupancy_type", "paired_offset_id",
        ],
    )
    return {
        "output_dir": tmp_path / "rapid_output",
        "summary": summary,
        "details": details,
        "bridges": bridges,
        "issues": issues,
        "reconciliation": reconciliation,
        "run_info": {
            "entity_name": "甲公司",
            "period_start": "2026-01-01",
            "period_end": "2026-12-31",
            "functional_currency": "CNY",
            "display_unit": "元",
            "display_divisor_minor": 100,
            "performance_materiality_minor": 50000,
            "protocol_version": "1.0",
            "run_status": "exported",
        },
        "decision_trace": decision_trace,
        "journal_snapshot": snapshot,
        "journal_sheet": "序时账",
        "journal_header_row": 1,
    }


def test_rapid_workbook_has_two_sheets(rapid_input):
    result = write_rapid_workbook(
        **{
            key: rapid_input[key]
            for key in (
                "summary", "details", "bridges", "issues",
                "reconciliation", "run_info", "decision_trace",
                "journal_snapshot", "journal_sheet", "journal_header_row",
            )
        },
        target=rapid_input["output_dir"] / "现金流量表附表计算结果_急速.xlsx",
    )
    workbook = openpyxl.load_workbook(result["path"], read_only=False)
    try:
        assert workbook.sheetnames == ["附表结果", "简报"]
    finally:
        workbook.close()


def test_rapid_mode_writes_sixteen_markdowns(rapid_input, rule_book):
    md_dir = rapid_input["output_dir"] / "底稿MD"
    result = write_item_markdowns(
        **{
            key: rapid_input[key]
            for key in (
                "summary", "details", "bridges", "issues",
                "decision_trace", "journal_snapshot", "journal_sheet",
                "journal_header_row",
            )
        },
        run_info=rapid_input["run_info"],
        output_dir=md_dir,
    )
    assert len(result["markdown_paths"]) == 16
    for item in rule_book.items:
        expected = md_dir / f"{item.item_no:02d}{item.name}.md"
        assert expected.is_file(), f"缺少 {expected.name}"
        content = expected.read_text(encoding="utf-8-sig")
        assert "审计底稿" in content
        assert "冻结金额" in content
    item1 = (md_dir / "01资产减值准备.md").read_text(encoding="utf-8-sig")
    assert "计提跌价准备" in item1
    assert "计入当期损益的减值" in item1
    assert "未命中" in item1
